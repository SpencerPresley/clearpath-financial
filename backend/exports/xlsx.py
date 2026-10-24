"""The items spreadsheet: one row per item, readable headers and values. Pure: rows in, xlsx bytes out."""

from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date
from typing import IO, Final

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from core.domain import Channel, ItemStatus, Market, Product, RiskTier
from exports.labels import CHANNEL_LABELS, PRODUCT_LABELS, RISK_LABELS, STATUS_LABELS


@dataclass(frozen=True)
class ExportRow:
    key: str
    title: str
    status: ItemStatus
    product: Product
    channel: Channel
    partner: str | None
    risk_tier: RiskTier
    round: int
    version_n: int | None
    markets: tuple[Market, ...]
    tags: tuple[str, ...]
    submitter: str | None
    reviewer: str | None
    open_findings: int
    submitted_on: date
    approved_on: date | None
    approval_expires_on: date | None
    legacy_ref: str | None


# (header, column width)
COLUMNS: Final = (
    ("Item", 10),
    ("Title", 48),
    ("Status", 18),
    ("Product", 16),
    ("Channel", 14),
    ("Partner", 20),
    ("Risk tier", 10),
    ("Round", 7),
    ("Version", 8),
    ("Markets", 12),
    ("Tags", 20),
    ("Submitter", 28),
    ("Reviewer", 20),
    ("Open findings", 13),
    ("Submitted", 12),
    ("Approved", 12),
    ("Approval expires", 16),
    ("Legacy Request ID", 18),
)
HEADERS: Final = tuple(h for h, _ in COLUMNS)
SHEET_TITLE: Final = "Items"
DATE_FORMAT: Final = "yyyy-mm-dd"
_DATE_COLUMNS: Final = frozenset({"Submitted", "Approved", "Approval expires"})


type CellValue = str | int | date | None


def _values(row: ExportRow) -> tuple[CellValue, ...]:
    return (
        row.key,
        row.title,
        STATUS_LABELS[row.status],
        PRODUCT_LABELS[row.product],
        CHANNEL_LABELS[row.channel],
        row.partner,
        RISK_LABELS[row.risk_tier],
        row.round,
        f"v{row.version_n}" if row.version_n is not None else None,
        ", ".join("National" if m is Market.NATIONAL else m.value for m in row.markets) or None,
        ", ".join(row.tags) or None,
        row.submitter,
        row.reviewer,
        row.open_findings,
        row.submitted_on,
        row.approved_on,
        row.approval_expires_on,
        row.legacy_ref,
    )


def write_items_xlsx(rows: Iterable[ExportRow], out: IO[bytes]) -> int:
    """Stream rows into a write-only workbook (constant memory for large exports). Returns the row count."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(SHEET_TITLE)
    ws.freeze_panes = "A2"
    for i, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    bold = Font(bold=True)
    header = []
    for name in HEADERS:
        cell = WriteOnlyCell(ws, value=name)
        cell.font = bold
        header.append(cell)
    ws.append(header)
    count = 0
    for row in rows:
        cells = []
        for name, value in zip(HEADERS, _values(row), strict=True):
            # openpyxl writes dates as dates; its stubs list only datetime.
            cell = WriteOnlyCell(ws, value=value)  # pyrefly: ignore[bad-argument-type]
            if name in _DATE_COLUMNS and value is not None:
                cell.number_format = DATE_FORMAT
            cells.append(cell)
        ws.append(cells)
        count += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{count + 1}"
    wb.save(out)
    return count
