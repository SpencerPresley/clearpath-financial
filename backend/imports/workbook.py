"""Read a tracker workbook once into typed, JSON-serializable data. Pure: openpyxl and pydantic, no Django.

openpyxl is loaded in normal mode, not read-only: read-only mode doesn't load `cell.hyperlink`, and the tracker
keeps file links in hyperlink targets whose display text is only a file name or a label.
"""

import io
import zipfile
from pathlib import Path
from collections.abc import Sequence
from datetime import UTC, date, datetime
from typing import IO, Final

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict

from imports.domain import ImportTargetField
from imports.mapping import guess_column_mapping

# The "Load sample tracker" workbook: a copy of samples/clearpath_review_tracker.xlsx (the Docker build context
# is backend/, so the repo-root copy isn't in the image). test_sample_copy keeps the two identical.
SAMPLE_TRACKER_PATH: Final = Path(__file__).resolve().parent / "sample" / "clearpath_review_tracker.xlsx"

MAX_BYTES: Final = 10 * 1024 * 1024
MAX_DATA_ROWS: Final = 5000
_HEADER_SCAN_ROWS: Final = 10
_MIN_HEADER_CELLS: Final = 3
_MIN_TRACKER_TARGETS: Final = 3


class UnreadableWorkbook(ValueError):
    """The upload isn't a workbook the importer can read. The message is one line for the admin."""


class _Frozen(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")


class RawCell(_Frozen):
    """One cell as typed in Excel. Exactly one of text/when/number is set unless the cell is empty. Strings
    are stripped; an all-space string is empty. Dates are UTC (openpyxl returns naive midnight)."""

    text: str | None = None
    when: datetime | None = None
    number: float | None = None
    hyperlink: str | None = None

    @property
    def is_empty(self) -> bool:
        return self.text is None and self.when is None and self.number is None

    @property
    def display(self) -> str | None:
        """The value as the admin would read it in the sheet."""
        if self.text is not None:
            return self.text
        if self.when is not None:
            return f"{self.when.month}/{self.when.day}/{self.when.year}"
        if self.number is not None:
            return str(int(self.number)) if self.number.is_integer() else str(self.number)
        return None


EMPTY: Final = RawCell()


class RawRow(_Frozen):
    row_number: int  # 1-based, as Excel shows it
    cells: tuple[RawCell, ...]  # one per header


class RawSheet(_Frozen):
    name: str
    header_row: int
    headers: tuple[str, ...]
    rows: tuple[RawRow, ...]  # every non-empty row below the header


class PartnerSheetRow(_Frozen):
    name: str
    type: str | None = None
    description: str | None = None
    contact_email: str | None = None
    agreement_date: date | None = None
    status: str | None = None
    notes: str | None = None


class RawWorkbook(_Frozen):
    tracker: RawSheet
    partners: tuple[PartnerSheetRow, ...] = ()
    reviewers: tuple[str, ...] = ()  # full names from a lists sheet's "Reviewer" column


def _cell(cell: Cell | MergedCell) -> RawCell:
    value = cell.value
    link = cell.hyperlink.target if isinstance(cell, Cell) and cell.hyperlink is not None else None
    link = link.strip() if link else None
    if value is None:
        return RawCell(hyperlink=link)
    if isinstance(value, bool):
        return RawCell(text="TRUE" if value else "FALSE", hyperlink=link)
    if isinstance(value, datetime):
        return RawCell(when=value.replace(tzinfo=UTC) if value.tzinfo is None else value, hyperlink=link)
    if isinstance(value, date):
        return RawCell(when=datetime(value.year, value.month, value.day, tzinfo=UTC), hyperlink=link)
    if isinstance(value, int | float):
        return RawCell(number=float(value), hyperlink=link)
    text = str(value).strip()
    return RawCell(text=text or None, hyperlink=link)


def _rows(ws: Worksheet) -> list[tuple[int, list[RawCell]]]:
    return [(cells[0].row or 0, [_cell(c) for c in cells]) for cells in ws.iter_rows() if cells]


def _header(rows: Sequence[tuple[int, list[RawCell]]]) -> tuple[int, list[str]] | None:
    for row_number, cells in rows[:_HEADER_SCAN_ROWS]:
        if sum(1 for c in cells if c.text is not None) >= _MIN_HEADER_CELLS:
            names = [c.display or f"Column {i + 1}" for i, c in enumerate(cells)]
            while names and names[-1].startswith("Column "):  # trailing blank header cells
                names.pop()
            return row_number, names
    return None


def _score(headers: Sequence[str]) -> int:
    return sum(1 for t in guess_column_mapping(headers).values() if t is not ImportTargetField.IGNORE)


def _sheet(name: str, rows: Sequence[tuple[int, list[RawCell]]], header_row: int, headers: list[str]) -> RawSheet:
    width = len(headers)
    body = []
    for row_number, cells in rows:
        if row_number <= header_row:
            continue
        padded = (cells + [EMPTY] * width)[:width]
        if all(c.is_empty for c in padded):
            continue
        body.append(RawRow(row_number=row_number, cells=tuple(padded)))
    if len(body) > MAX_DATA_ROWS:
        raise UnreadableWorkbook(f"The sheet has {len(body)} rows; split it into files of {MAX_DATA_ROWS} or fewer.")
    return RawSheet(name=name, header_row=header_row, headers=tuple(headers), rows=tuple(body))


def _column(headers: Sequence[str], *names: str) -> int | None:
    folded = [h.casefold() for h in headers]
    for name in names:
        if name in folded:
            return folded.index(name)
    return None


def _partners(sheet: RawSheet) -> tuple[PartnerSheetRow, ...]:
    h = sheet.headers
    name_i = _column(h, "partner name", "name")
    if name_i is None:
        return ()
    cols = {
        "type": _column(h, "type"),
        "description": _column(h, "description"),
        "contact_email": _column(h, "contact email", "email"),
        "agreement_date": _column(h, "agreement date", "agreement signed"),
        "status": _column(h, "status"),
        "notes": _column(h, "notes"),
    }
    out = []
    for row in sheet.rows:
        name = row.cells[name_i].display
        if name is None:
            continue

        def text(key: str, row: RawRow = row) -> str | None:
            i = cols[key]
            return row.cells[i].display if i is not None else None

        agreed_i = cols["agreement_date"]
        agreed = row.cells[agreed_i].when if agreed_i is not None else None
        out.append(
            PartnerSheetRow(
                name=name,
                type=text("type"),
                description=text("description"),
                contact_email=text("contact_email"),
                agreement_date=agreed.date() if agreed else None,
                status=text("status"),
                notes=text("notes"),
            )
        )
    return tuple(out)


def _reviewers(sheet: RawSheet) -> tuple[str, ...]:
    i = _column(sheet.headers, "reviewer", "reviewers")
    if i is None:
        return ()
    return tuple(name for row in sheet.rows if (name := row.cells[i].display))


def read_workbook(source: IO[bytes]) -> RawWorkbook:
    """Parse an uploaded tracker. The tracker sheet is the one whose headers match the most import targets;
    a sheet with "Partner Name" + "Type" supplies the partner directory; a "Reviewer" column supplies names."""
    data = source.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise UnreadableWorkbook("The file is over 10 MB; remove unused sheets or split it.")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise UnreadableWorkbook("This file isn't an .xlsx workbook. Save it from Excel as .xlsx and retry.") from exc

    sheets: list[tuple[str, list[tuple[int, list[RawCell]]], int, list[str]]] = []
    for ws in wb.worksheets:
        rows = _rows(ws)
        header = _header(rows)
        if header is not None:
            sheets.append((ws.title, rows, header[0], header[1]))
    scored = sorted(((_score(s[3]), i) for i, s in enumerate(sheets)), key=lambda p: (-p[0], p[1]))
    if not scored or scored[0][0] < _MIN_TRACKER_TARGETS:
        raise UnreadableWorkbook("No sheet has tracker columns (Request ID, Status, Product, ...) in its header.")
    tracker_i = scored[0][1]
    tracker = _sheet(*sheets[tracker_i])

    partners: tuple[PartnerSheetRow, ...] = ()
    reviewers: tuple[str, ...] = ()
    for i, (name, rows, header_row, headers) in enumerate(sheets):
        if i == tracker_i:
            continue
        folded = {h.casefold() for h in headers}
        if not partners and "type" in folded and folded & {"partner name", "name"}:
            partners = _partners(_sheet(name, rows, header_row, headers))
        elif not reviewers and folded & {"reviewer", "reviewers"}:
            reviewers = _reviewers(_sheet(name, rows, header_row, headers))
    return RawWorkbook(tracker=tracker, partners=partners, reviewers=reviewers)
