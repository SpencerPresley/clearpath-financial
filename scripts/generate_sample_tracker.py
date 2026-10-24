#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.13"
# dependencies = ["openpyxl>=3.1.5"]
# ///
"""Generate the sample legacy Excel review tracker and its import oracle.

Writes two files:
  samples/clearpath_review_tracker.xlsx  the messy spreadsheet the Import page ingests
  samples/README.md                      what's in it, including every intended inconsistency

Run from the repo root:  uv run scripts/generate_sample_tracker.py

Everything is driven by one seeded random.Random, so re-running produces identical cell values.
Edit this script, not the outputs; the README's oracle tables are generated from the same data.
"""

import random
import sys
from collections import Counter
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from enum import IntEnum, StrEnum
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------------------------
# Knobs
# ---------------------------------------------------------------------------------------------

SEED = 20250908
# The tracker ends the Friday before ClearPath Compliance Review went live (Monday 2026-08-17), so the imported
# history and the demo seed's native history meet without overlapping. The seed's go-live is
# backend/demo/content/calendar.py `GO_LIVE`; demo/tests/test_seed.py fails if the two ever overlap.
AS_OF = date(2026, 8, 14)  # the day the team last touched the tracker (a Friday)
FIRST_SUBMISSION = date(2025, 9, 2)  # the Tuesday after Labor Day
LAST_SUBMISSION = date(2026, 8, 12)
FREEZE_START = date(2025, 12, 22)  # holiday review freeze: nothing reviewed in this window
FREEZE_END = date(2026, 1, 2)
RATE_SHEET_CHANGE = date(2026, 6, 1)  # PL APR floor moved 8.99% -> 9.49%
REG_B_REWRITE = date(2026, 7, 21)
NEW_REVIEWER_START = date(2026, 3, 2)

# (year, month, submissions). Volume roughly doubles over the year (December is the holiday freeze); August
# stops at the 12th, the rush before the cutover.
MONTHLY_VOLUME: tuple[tuple[int, int, int], ...] = (
    (2025, 9, 9),
    (2025, 10, 11),
    (2025, 11, 12),
    (2025, 12, 9),
    (2026, 1, 13),
    (2026, 2, 15),
    (2026, 3, 16),
    (2026, 4, 17),
    (2026, 5, 19),
    (2026, 6, 20),
    (2026, 7, 23),
    (2026, 8, 14),
)

HOLIDAYS = frozenset(
    {
        date(2025, 9, 1),
        date(2025, 11, 27),
        date(2025, 11, 28),
        date(2025, 12, 24),
        date(2025, 12, 25),
        date(2025, 12, 26),
        date(2025, 12, 31),
        date(2026, 1, 1),
        date(2026, 1, 19),
        date(2026, 2, 16),
        date(2026, 5, 25),
        date(2026, 7, 3),
        date(2026, 9, 7),
    }
)
SATURDAY = 5

# Review timing, in business days. The first-review wait grows with volume; that is the pain.
FIRST_WAIT_BASE = 2.0
FIRST_WAIT_GROWTH_PER_MONTH = 0.5
FIRST_WAIT_JITTER = 3
LATER_WAIT_BASE = 1.0
LATER_WAIT_GROWTH_PER_MONTH = 0.3
LATER_WAIT_JITTER = 2
HIGH_RISK_EXTRA_WAIT = 1
SPONSOR_BANK_EXTRA_WAIT = 4
MARKETER_FAST = (1, 3)
MARKETER_SLOW = (5, 9)
MARKETER_SLOW_P = 0.15
ASSIGN_DELAY_MAX = 2

MAX_ROUNDS = 5
RESOLVE_P = 0.45
NEW_ISSUE_P = 0.25
NEW_ISSUE_MAX_ROUND = 4
REJECT_P = 0.35
REVIEWER_FIXES_P = 0.35
ABANDON_P = 0.05
ABANDON_MIN_AGE_DAYS = 75
HANDOFF_P = 0.06
PRE_REVISED_P = 0.08
UNASSIGNED_WINDOW_DAYS = 10
UNASSIGNED_P = 0.7
LOW_RISK_THEME_P = 0.15
TRIGGER_WITHOUT_FINDING_P = 0.3
RE_REVIEW_P = 0.25
RE_REVIEW_WINDOW = (date(2026, 5, 1), date(2026, 8, 12))
RE_REVIEW_LOOKAHEAD_DAYS = 30

# Placement bookkeeping
LAUNCH_DELAY = (1, 5)
SCREENSHOT_P = 0.7
PLACEMENT_UNRECORDED_P = 0.1
SEND_WINDOW_DAYS = (14, 45)
MONTHS_EVERGREEN = 12
MONTHS_PARTNER = 6

# Link rendering
HYPERLINK_ERA = date(2026, 2, 1)  # the team started pasting real links around February
HYPERLINK_P_EARLY = 0.2
HYPERLINK_P_LATE = 0.65
NAME_ONLY_P = 0.3
FRIENDLY_LINK_TEXT_P = 0.3

# Layout
DUPLICATE_SOURCE_INDEX = 61  # 0-based item index whose row is pasted twice
DUPLICATE_OFFSET = 3  # the stray copy lands this many rows after the original
REUSED_ID_INDEX = 118  # this item was copied from the row above and kept its Request ID

# Messiness volumes (cells per variant)
STATUS_VARIANTS_PER_KIND = 3
APPROVED_TRAILING_SPACE = 4
MINOR_EDITS_COUNT = 4
LOWERCASE_AWC_COUNT = 3
PRODUCT_ABBREV_PL = 14
PRODUCT_LOWER_PL = 4
PRODUCT_ABBREV_CC = 10
PRODUCT_ABBREV_MTG = 6
DATE_SLASH_TEXT = 6
DATE_ISO_TEXT = 5
DATE_NO_YEAR = 4
EXPIRY_TEXT = 3
APPROVER_ONLY = 6
BLANK_EXPIRY = 6
YN_YES = 5
YN_BLANK = 5
YN_UNKNOWN = 3

OUT_DIR = Path(__file__).resolve().parents[1] / "samples"
XLSX_PATH = OUT_DIR / "clearpath_review_tracker.xlsx"
README_PATH = OUT_DIR / "README.md"
SHAREPOINT = "https://sharepoint.clearpath.example/sites/MarketingCompliance"
CLEARPATH_SITE = "https://www.clearpathfinancial.example"
FILE_CREATED = datetime(2025, 8, 28, 16, 12, tzinfo=UTC)

# ---------------------------------------------------------------------------------------------
# Domain vocabulary
# ---------------------------------------------------------------------------------------------


# Enum values for Product, ItemStatus, Outcome and PartnerKind match backend/core/domain.py, so the README's
# "canonical" column is exactly what the importer should store. Channel and Risk values are the sheet labels;
# their domain values are derived in DOMAIN_* below.


class Product(StrEnum):
    PL = "personal_loan"
    CC = "credit_card"
    MTG = "mortgage_prequal"

    @property
    def code(self) -> str:
        return self.name


class Channel(StrEnum):
    EMAIL = "Email"
    PAID_SOCIAL = "Paid Social"
    PAID_SEARCH = "Paid Search"
    DIRECT_MAIL = "Direct Mail"
    AFFILIATE_PAGE = "Affiliate Page"
    SMS = "SMS"
    LANDING_PAGE = "Landing Page"
    CALL_SCRIPT = "Call Script"
    DISPLAY = "Display"


class Risk(StrEnum):
    HIGH = "High"
    MEDIUM = "Medium"
    LOW = "Low"


class ItemStatus(StrEnum):
    SUBMITTED = "submitted"
    IN_REVIEW = "in_review"
    CHANGES_REQUESTED = "changes_requested"
    APPROVED = "approved"
    REJECTED = "rejected"


class Outcome(StrEnum):
    APPROVED = "approved"
    APPROVED_WITH_EDITS = "approved_with_edits"
    CHANGES_REQUESTED = "changes_requested"
    REJECTED = "rejected"


class PartnerKind(StrEnum):
    AFFILIATE = "affiliate"
    LEAD_GEN = "lead_gen"
    MARKETPLACE = "marketplace"
    INFLUENCER = "influencer"


class Col(IntEnum):
    REQUEST_ID = 0
    DATE_SUBMITTED = 1
    SUBMITTER = 2
    PRODUCT = 3
    CHANNEL = 4
    ASSET = 5
    PARTNER = 6
    RISK = 7
    TRIGGER = 8
    REVIEWER = 9
    VERSION = 10
    ROUND = 11
    STATUS = 12
    COMMENTS = 13
    APPROVAL = 14
    EXPIRATION = 15
    LIVE = 16


HEADERS: tuple[str, ...] = (
    "Request ID",
    "Date Submitted",
    "Submitter / BU",
    "Product",
    "Channel",
    "Asset Name / File Link",
    "Affiliate / Partner",
    "Risk Tier",
    "Trigger Terms Present (Y/N)",
    "Assigned Reviewer",
    "Version #",
    "Revision Round",
    "Status",
    "Reviewer Comments / Required Changes",
    "Approval Date & Approver",
    "Approval Expiration / Re-review Date",
    "Live URL / Placement & Screenshot Link",
)
COLUMN_WIDTHS: tuple[int, ...] = (13, 14, 34, 17, 14, 52, 20, 10, 12, 16, 9, 10, 22, 90, 20, 18, 60)
LETTERS = "ABCDEFGHIJKLMNOPQ"

PRODUCT_LABEL = {Product.PL: "Personal Loan", Product.CC: "Credit Card", Product.MTG: "Mortgage Prequal"}
STATUS_LIST: tuple[str, ...] = (
    "Not Started",
    "In Review",
    "Changes Requested",
    "Approved",
    "Approved w/ Changes",
    "Rejected",
)
STATUS_DEFAULT_LABEL: dict[tuple[ItemStatus, Outcome | None], str] = {
    (ItemStatus.SUBMITTED, None): "Not Started",
    (ItemStatus.IN_REVIEW, None): "In Review",
    (ItemStatus.CHANGES_REQUESTED, None): "Changes Requested",
    (ItemStatus.APPROVED, Outcome.APPROVED): "Approved",
    (ItemStatus.APPROVED, Outcome.APPROVED_WITH_EDITS): "Approved w/ Changes",
    (ItemStatus.REJECTED, Outcome.REJECTED): "Rejected",
}
PARTNER_KIND_LABEL = {
    PartnerKind.AFFILIATE: "Affiliate",
    PartnerKind.LEAD_GEN: "Lead Gen",
    PartnerKind.MARKETPLACE: "Marketplace",
    PartnerKind.INFLUENCER: "Influencer",
}
MONTH_ABBR = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
ONE_OFF_CHANNELS = frozenset({Channel.EMAIL, Channel.SMS, Channel.DIRECT_MAIL})
EVERGREEN_CHANNELS = frozenset({Channel.LANDING_PAGE, Channel.CALL_SCRIPT})
WEB_CHANNELS = frozenset({Channel.AFFILIATE_PAGE, Channel.LANDING_PAGE})
MEDIUM_RISK_CHANNELS = frozenset({Channel.PAID_SOCIAL, Channel.PAID_SEARCH, Channel.DISPLAY, Channel.LANDING_PAGE})


@dataclass(frozen=True)
class Marketer:
    name: str
    initials: str
    bu: str
    products: tuple[Product, ...]
    channels: tuple[Channel, ...]
    weight_start: float
    weight_end: float
    partner_led: bool = False

    @property
    def submitter_cell(self) -> str:
        return f"{self.name} / {self.bu}"


@dataclass(frozen=True)
class Reviewer:
    name: str
    initials: str
    active_from: date
    weight: float
    is_lead: bool = False

    @property
    def approver_label(self) -> str:
        first, last = self.name.split(" ", 1)
        return f"{first[0]}. {last}"


@dataclass(frozen=True)
class Partner:
    name: str
    kind: PartnerKind
    description: str
    contact_email: str
    agreement_date: date
    status: str
    notes: str
    domain: str
    path: str
    products: tuple[Product, ...]
    channels: tuple[Channel, ...]
    weight: float


@dataclass(frozen=True)
class Theme:
    name: str
    product: Product
    low_risk: bool = False
    months: tuple[int, ...] = ()


@dataclass(frozen=True)
class Issue:
    """One compliance finding a reviewer can raise, phrased the way it lands in the comment cell."""

    key: str
    flag: str
    label: str
    still: str
    products: tuple[Product, ...]
    channels: tuple[Channel, ...] = ()  # empty means any channel
    generic: bool = False  # the only kind raised on low-risk (no-offer) assets
    trigger: bool = False  # the copy contains a Reg Z trigger term
    partner_only: bool = False
    blocking: bool = False  # can get the asset rejected outright
    minor: bool = False  # small enough that a reviewer may fix it directly
    sponsor_bank: bool = False  # adds the sponsor bank's turnaround
    on_or_after: date | None = None
    before: date | None = None


@dataclass(frozen=True)
class Terms:
    pl_floor: str
    pl_ceil: str
    cc_low: str
    cc_high: str
    mtg_rate: str
    mtg_apr: str


MARKETERS: tuple[Marketer, ...] = (
    Marketer(
        "Priya Natarajan",
        "PN",
        "Growth Marketing (PL)",
        (Product.PL,),
        (
            Channel.LANDING_PAGE,
            Channel.PAID_SEARCH,
            Channel.DISPLAY,
            Channel.DIRECT_MAIL,
            Channel.EMAIL,
            Channel.CALL_SCRIPT,
        ),
        0.24,
        0.20,
    ),
    Marketer(
        "Marcus Bell",
        "MB",
        "Card Marketing",
        (Product.CC,),
        (Channel.EMAIL, Channel.DIRECT_MAIL, Channel.PAID_SOCIAL, Channel.LANDING_PAGE, Channel.DISPLAY),
        0.20,
        0.18,
    ),
    Marketer(
        "Elena Ruiz",
        "ER",
        "Mortgage Marketing",
        (Product.MTG,),
        (Channel.LANDING_PAGE, Channel.PAID_SEARCH, Channel.CALL_SCRIPT, Channel.EMAIL, Channel.PAID_SOCIAL),
        0.14,
        0.14,
    ),
    Marketer(
        "Jordan Whitfield",
        "JW",
        "Partnerships & Affiliates",
        (Product.PL, Product.CC, Product.MTG),
        (),
        0.14,
        0.26,
        partner_led=True,
    ),
    Marketer(
        "Hannah Cho",
        "HC",
        "Lifecycle & CRM",
        (Product.PL, Product.CC),
        (Channel.EMAIL, Channel.SMS),
        0.16,
        0.12,
    ),
    Marketer(
        "Andre Duval",
        "AD",
        "Brand & Paid Media",
        (Product.PL, Product.CC, Product.MTG),
        (Channel.PAID_SOCIAL, Channel.DISPLAY, Channel.PAID_SEARCH),
        0.12,
        0.10,
    ),
)

REVIEWERS: tuple[Reviewer, ...] = (
    Reviewer("Dana Kim", "DK", FIRST_SUBMISSION, 1.0),
    Reviewer("Samuel Ortega", "SO", FIRST_SUBMISSION, 1.0),
    Reviewer("Leah Brandt", "LB", FIRST_SUBMISSION, 0.9),
    Reviewer("Nikhil Rao", "NR", NEW_REVIEWER_START, 1.0),
    Reviewer("Rachel Okafor", "RO", FIRST_SUBMISSION, 0.35, is_lead=True),
)

PARTNERS: tuple[Partner, ...] = (
    Partner(
        "LendCompare",
        PartnerKind.MARKETPLACE,
        "Personal loan and card comparison site",
        "partnerships@lendcompare.example",
        date(2024, 3, 18),
        "Active",
        "Top PL volume partner. Quarterly screenshot audit.",
        "lendcompare.example",
        "personal-loans/clearpath",
        (Product.PL, Product.CC),
        (Channel.AFFILIATE_PAGE, Channel.EMAIL),
        1.4,
    ),
    Partner(
        "RateRadar",
        PartnerKind.MARKETPLACE,
        "Loan rate comparison site",
        "affiliates@rateradar.example",
        date(2024, 9, 9),
        "Active",
        "",
        "rateradar.example",
        "lenders/clearpath-financial",
        (Product.PL,),
        (Channel.AFFILIATE_PAGE,),
        0.8,
    ),
    Partner(
        "CardPerks Weekly",
        PartnerKind.AFFILIATE,
        "Credit card rewards blog and newsletter",
        "editor@cardperksweekly.example",
        date(2025, 2, 3),
        "Active",
        "Newsletter sends need our review too, not just the site.",
        "cardperksweekly.example",
        "reviews/clearpath-cash-back-card",
        (Product.CC,),
        (Channel.AFFILIATE_PAGE, Channel.EMAIL),
        1.0,
    ),
    Partner(
        "Money with Maya",
        PartnerKind.INFLUENCER,
        "Personal-finance influencer (IG / TikTok @moneywithmaya)",
        "maya@moneywithmaya.example",
        date(2025, 8, 21),
        "Active",
        "Warning letter 3/2026 re: missing #ad on two posts.",
        "moneywithmaya.example",
        "links/clearpath",
        (Product.PL, Product.CC),
        (Channel.PAID_SOCIAL,),
        0.9,
    ),
    Partner(
        "LeadStream Direct",
        PartnerKind.LEAD_GEN,
        "Lead aggregator (PL), email + SMS to its own lists",
        "compliance@leadstreamdirect.example",
        date(2024, 6, 12),
        "Active",
        "Commission hold 7/2026-8/2026 after an unapproved SMS went out.",
        "leadstreamdirect.example",
        "offers/clearpath",
        (Product.PL,),
        (Channel.EMAIL, Channel.SMS),
        0.9,
    ),
    Partner(
        "FinMatch",
        PartnerKind.MARKETPLACE,
        "Loan marketplace with soft-pull prequal",
        "partners@finmatch.example",
        date(2025, 5, 14),
        "Active",
        "",
        "finmatch.example",
        "marketplace/clearpath",
        (Product.PL,),
        (Channel.AFFILIATE_PAGE,),
        0.8,
    ),
    Partner(
        "HomeBound Leads",
        PartnerKind.LEAD_GEN,
        "Mortgage lead generation (CA), co-branded LPs and warm transfers",
        "ops@homeboundleads.example",
        date(2025, 10, 1),
        "Active",
        "Warm-transfer scripts go through us before every refresh.",
        "homeboundleads.example",
        "ca/clearpath-prequal",
        (Product.MTG,),
        (Channel.LANDING_PAGE, Channel.CALL_SCRIPT),
        0.7,
    ),
    Partner(
        "Golden State Saver",
        PartnerKind.AFFILIATE,
        "California personal-finance blog",
        "hello@goldenstatesaver.example",
        date(2026, 1, 15),
        "Active",
        "Onboarded 1/2026; sent the disclosure kit.",
        "goldenstatesaver.example",
        "best-of/clearpath",
        (Product.PL, Product.CC, Product.MTG),
        (Channel.AFFILIATE_PAGE, Channel.EMAIL),
        0.8,
    ),
)

THEMES: tuple[Theme, ...] = (
    Theme("Debt Consolidation", Product.PL),
    Theme("Home Improvement", Product.PL, months=(3, 4, 5, 6, 7, 8, 9)),
    Theme("Tax Season", Product.PL, months=(1, 2, 3, 4)),
    Theme("Credit Card Payoff", Product.PL),
    Theme("Wedding Season", Product.PL, months=(3, 4, 5, 6)),
    Theme("Medical Bills", Product.PL),
    Theme("Back to School", Product.PL, months=(7, 8)),
    Theme("Holiday Expenses", Product.PL, months=(11, 12)),
    Theme("Check Your Rate", Product.PL),
    Theme("Customer Story", Product.PL, low_risk=True),
    Theme("Brand Awareness", Product.PL, low_risk=True),
    Theme("Cash Back Card", Product.CC),
    Theme("0% Intro APR", Product.CC),
    Theme("No Annual Fee", Product.CC),
    Theme("Double Points", Product.CC, months=(5, 6, 7, 8)),
    Theme("Balance Transfer", Product.CC),
    Theme("Holiday Spend", Product.CC, months=(11, 12)),
    Theme("Card Brand", Product.CC, low_risk=True),
    Theme("First-Time Buyer Prequal", Product.MTG),
    Theme("Spring Homebuying", Product.MTG, months=(2, 3, 4, 5)),
    Theme("Rate Check", Product.MTG),
    Theme("Bay Area Buyers", Product.MTG),
    Theme("Prequal in Minutes", Product.MTG),
    Theme("Homebuyer Tips", Product.MTG, low_risk=True),
)

CHANNEL_NOUNS: dict[Channel, tuple[str, ...]] = {
    Channel.EMAIL: ("email", "email blast", "nurture email"),
    Channel.PAID_SOCIAL: ("social ad", "carousel", "IG Reel"),
    Channel.PAID_SEARCH: ("search ads", "RSA copy"),
    Channel.DIRECT_MAIL: ("mailer", "postcard", "letter package"),
    Channel.AFFILIATE_PAGE: ("listing", "review page", "offer module"),
    Channel.SMS: ("SMS", "text blast"),
    Channel.LANDING_PAGE: ("landing page", "LP"),
    Channel.CALL_SCRIPT: ("call script", "warm-transfer script"),
    Channel.DISPLAY: ("display banners", "programmatic banners"),
}
CHANNEL_CODE = {
    Channel.EMAIL: ("EM", "pdf"),
    Channel.PAID_SOCIAL: ("SOC", "png"),
    Channel.PAID_SEARCH: ("SEM", "xlsx"),
    Channel.DIRECT_MAIL: ("DM", "pdf"),
    Channel.AFFILIATE_PAGE: ("AFF", "pdf"),
    Channel.SMS: ("SMS", "docx"),
    Channel.LANDING_PAGE: ("LP", "pdf"),
    Channel.CALL_SCRIPT: ("SCRIPT", "docx"),
    Channel.DISPLAY: ("DSP", "zip"),
}
SEGMENTS = {
    Product.PL: ("PL prospects seg B", "PL abandoned-app list", "PL past customers"),
    Product.CC: ("CC prospects (prescreen file)", "cardholders - spend stim", "CC waitlist"),
    Product.MTG: ("MTG prequal leads 30-60d", "MTG newsletter list"),
}

ALL = (Product.PL, Product.CC, Product.MTG)
PL_CC = (Product.PL, Product.CC)
PARTNER_CHANNELS = (Channel.AFFILIATE_PAGE, Channel.PAID_SOCIAL, Channel.EMAIL)
NOT_SEARCH = tuple(c for c in Channel if c != Channel.PAID_SEARCH)

ISSUES: tuple[Issue, ...] = (
    Issue(
        "as_low_as",
        "'as low as {pl_floor}% APR' needs the full range ({pl_floor}%-{pl_ceil}% APR) "
        "+ 'based on creditworthiness' qualifier",
        "range",
        "rate still missing range + creditworthiness qualifier",
        (Product.PL,),
        NOT_SEARCH,
    ),
    Issue(
        "pl_payment",
        "'payments from $199/mo' is a trigger term - add APR, term and total of payments; "
        "use the rep example ($10,000 / 36 mo / 12.99% APR = $337/mo, $12,128 total)",
        "rep example",
        "rep example still missing total of payments",
        (Product.PL,),
        NOT_SEARCH,
        trigger=True,
    ),
    Issue(
        "pl_term",
        "'pay it off in 36 months' is a trigger term - needs APR + full repayment terms",
        "repayment terms",
        "36-month claim still has no APR next to it",
        (Product.PL,),
        trigger=True,
    ),
    Issue(
        "cfl_license",
        "add CA license line ('Loans made or arranged pursuant to a California Financing Law license, #60DBO-184027')",
        "license line",
        "CFL license line still missing",
        (Product.PL,),
        NOT_SEARCH,
    ),
    Issue(
        "soft_pull",
        "'won't affect your credit score' - limit it to checking your rate (soft pull); applying is a hard inquiry",
        "soft-pull wording",
        "credit-score line still not limited to the rate check",
        (Product.PL, Product.MTG),
    ),
    Issue(
        "guaranteed",
        "'guaranteed approval' - remove, we can't promise approval",
        "'guaranteed' removed",
        "'guaranteed' still in the subhead",
        PL_CC,
        blocking=True,
    ),
    Issue(
        "preapproved",
        "'You're pre-approved!' - this is not a prescreened firm offer; use 'check your rate'",
        "'pre-approved' removed",
        "'pre-approved' still in the subject line / headline",
        PL_CC,
        (Channel.EMAIL, Channel.DIRECT_MAIL, Channel.SMS, Channel.AFFILIATE_PAGE, Channel.PAID_SOCIAL),
        blocking=True,
    ),
    Issue(
        "apr_prominence",
        "'interest rate' in the headline is bigger than the APR - APR must be at least as prominent",
        "APR prominence",
        "APR still smaller than 'interest rate'",
        (Product.PL,),
        minor=True,
    ),
    Issue(
        "no_fees",
        "'no fees' isn't accurate (0%-8% origination fee) - say 'no prepayment penalty' instead",
        "fee claim",
        "'no fees' still in body copy",
        (Product.PL,),
        minor=True,
    ),
    Issue(
        "fast_funding",
        "'money in your account today' - need ops substantiation or change to 'as soon as next business day'",
        "funding claim",
        "same-day funding claim still unsupported",
        (Product.PL,),
        minor=True,
    ),
    Issue(
        "stale_floor",
        "'8.99%' is the old floor - rate sheet moved to 9.49% APR eff 6/1",
        "rate update",
        "still showing the old 8.99% floor",
        (Product.PL,),
        on_or_after=RATE_SHEET_CHANGE,
        before=date(2026, 8, 15),
    ),
    Issue(
        "fixed_pl",
        "'fixed rate' is fine (PL is fixed) but drop 'locked in forever'",
        "'forever' dropped",
        "'locked in forever' still there",
        (Product.PL,),
        minor=True,
    ),
    Issue(
        "cc_apr",
        "rate shown without 'APR' or 'variable' - add '{cc_low}%-{cc_high}% variable APR'",
        "APR line",
        "variable APR line still missing",
        (Product.CC,),
        trigger=True,
    ),
    Issue(
        "cc_intro",
        "'0% for 15 months' - needs 'intro' in the same phrase, the end of the intro period and the go-to APR "
        "({cc_low}%-{cc_high}% variable)",
        "intro APR wording",
        "go-to APR still not near the first mention of 0%",
        (Product.CC,),
        trigger=True,
    ),
    Issue(
        "cc_no_fee",
        "'no annual fee' triggers the full disclosure - add the APR (variable) and other fees",
        "fee disclosures",
        "'no annual fee' still has no APR next to it",
        (Product.CC,),
        trigger=True,
    ),
    Issue(
        "cc_issuer",
        "missing issuer line ('The ClearPath Card is issued by Harborview Bank, N.A., Member FDIC')",
        "issuer line",
        "issuer line still missing from footer",
        (Product.CC,),
        NOT_SEARCH,
        minor=True,
    ),
    Issue(
        "cc_sponsor",
        "Harborview (sponsor bank) has to sign off on new card creative - sent to them, holding until they reply",
        "Harborview sign-off",
        "still waiting on Harborview",
        (Product.CC,),
        sponsor_bank=True,
    ),
    Issue(
        "cc_rewards",
        "'unlimited 2% cash back' - add 'on purchases' + link to the rewards terms",
        "rewards terms",
        "rewards terms link still missing",
        (Product.CC,),
        minor=True,
    ),
    Issue(
        "prescreen",
        "'You've been pre-selected' - needs the prescreen short notice (front page, bordered, 12pt) + "
        "PRESCREEN & OPT-OUT NOTICE on the back; send me the selection criteria for the file",
        "prescreen notices",
        "long notice still under 8pt / not set apart",
        PL_CC,
        (Channel.DIRECT_MAIL, Channel.EMAIL),
    ),
    Issue(
        "nmls",
        "add 'NMLS ID #1984302' + www.nmlsconsumeraccess.org",
        "NMLS line",
        "NMLS line still missing",
        (Product.MTG,),
        minor=True,
    ),
    Issue(
        "mtg_govt",
        "eagle icon + 'government-backed' wording - we're not FHA-approved, remove both",
        "gov't imagery removed",
        "eagle icon still in the hero",
        (Product.MTG,),
        blocking=True,
    ),
    Issue(
        "mtg_payment",
        "'$2,140/mo' needs every payment stream + 'does not include taxes and insurance; actual payment will be "
        "greater' right next to it, not in a footnote",
        "payment disclosure",
        "taxes & insurance line still in the footnote",
        (Product.MTG,),
        trigger=True,
    ),
    Issue(
        "mtg_arm",
        "'fixed' used for the 7/6 ARM - 'ARM' has to come first, with the fixed period",
        "ARM wording",
        "'fixed' still precedes 'ARM'",
        (Product.MTG,),
    ),
    Issue(
        "mtg_preapproved",
        "'get pre-approved in minutes' - this is prequal only, say 'see if you prequalify'",
        "prequal wording",
        "'pre-approved' still in the CTA",
        (Product.MTG,),
    ),
    Issue(
        "mtg_rate",
        "'{mtg_rate}%' needs the APR alongside ('{mtg_rate}% rate / {mtg_apr}% APR') at equal prominence",
        "rate/APR pairing",
        "APR still smaller than the rate",
        (Product.MTG,),
    ),
    Issue(
        "crmla",
        "add the CRMLA license line (licensed by the DFPI under the California Residential Mortgage Lending Act)",
        "CRMLA line",
        "CRMLA line still missing",
        (Product.MTG,),
        NOT_SEARCH,
        minor=True,
    ),
    Issue(
        "canspam",
        "footer is missing our postal address + unsubscribe link",
        "footer",
        "unsubscribe link still missing",
        ALL,
        (Channel.EMAIL,),
        minor=True,
    ),
    Issue(
        "email_subject",
        "subject line 'Your approval is waiting' doesn't match the body - nobody's approved yet",
        "subject line",
        "subject line still implies an approval",
        ALL,
        (Channel.EMAIL,),
    ),
    Issue(
        "sms_consent",
        "need the consent record (with 'not a condition of purchase'), 'Reply STOP to opt out' in the copy, "
        "and the send window inside 8am-9pm recipient time",
        "consent/STOP",
        "STOP language still missing",
        PL_CC,
        (Channel.SMS,),
    ),
    Issue(
        "sms_rate",
        "'rates from {pl_floor}%' in a text needs the range + qualifier, or drop the rate",
        "SMS rate claim",
        "SMS rate still has no qualifier",
        (Product.PL,),
        (Channel.SMS,),
    ),
    Issue(
        "partner_disclosure",
        "partner copy has no '#ad' / paid-partner disclosure",
        "#ad disclosure",
        "'#ad' still below the fold",
        ALL,
        PARTNER_CHANNELS,
        partner_only=True,
    ),
    Issue(
        "partner_old_rates",
        "partner page is quoting last quarter's rates - update to the current rate sheet",
        "rates updated",
        "partner page still on the old rates",
        ALL,
        (Channel.AFFILIATE_PAGE,),
        partner_only=True,
    ),
    Issue(
        "partner_best",
        "partner headline calls us 'the best in California' - we can't substantiate 'best', they need to drop it",
        "'best' claim",
        "'best' still in the partner headline",
        ALL,
        PARTNER_CHANNELS,
        partner_only=True,
    ),
    Issue(
        "reg_b",
        "imagery + zip-code targeting list - flagging for fair lending review (Reg B 1002.4(b))",
        "fair lending review",
        "fair lending review still open on the targeting list",
        ALL,
        (Channel.PAID_SOCIAL, Channel.DISPLAY, Channel.DIRECT_MAIL),
        before=REG_B_REWRITE,
    ),
    Issue(
        "reg_b_2026",
        "imagery + zip-code targeting list - flagging for fair lending review under the rewritten "
        "Reg B 1002.4(b) (eff 7/21/26)",
        "fair lending review",
        "fair lending review still open on the targeting list",
        ALL,
        (Channel.PAID_SOCIAL, Channel.DISPLAY, Channel.DIRECT_MAIL),
        on_or_after=REG_B_REWRITE,
    ),
    Issue(
        "search_rate",
        "ad text has the rate but no room for range + qualifier - drop the rate from the ad (LP still needs it)",
        "rate dropped",
        "headline 2 still has the rate",
        ALL,
        (Channel.PAID_SEARCH,),
    ),
    Issue(
        "call_soft",
        "script says 'this won't touch your credit' right before the hard-pull consent - keep it to the prequal step",
        "soft-pull scope",
        "credit line still sits before the hard-pull consent",
        (Product.PL, Product.MTG),
        (Channel.CALL_SCRIPT,),
    ),
    Issue(
        "call_apr",
        "agent has to read the APR range verbatim, not 'rates start low'",
        "APR read-out",
        "APR range still paraphrased",
        (Product.PL, Product.MTG),
        (Channel.CALL_SCRIPT,),
    ),
    Issue(
        "superlative",
        "'fastest loans in California' - substantiate or soften",
        "superlative",
        "'fastest' still in the headline",
        ALL,
        generic=True,
        minor=True,
    ),
    Issue(
        "testimonial",
        "testimonial needs a typical-results disclosure or has to go",
        "testimonial disclosure",
        "testimonial disclosure still missing",
        ALL,
        generic=True,
        minor=True,
    ),
    Issue(
        "logo",
        "not a compliance issue, but the logo lockup is the old one",
        "logo",
        "old logo still there",
        ALL,
        generic=True,
        minor=True,
    ),
)

APPROVE_PHRASES = ("approved.", "approved, ok to publish.", "good to go - approved.", "looks good, approved.")
NOTE_ROW_TEXT = (
    "NOTE (R. Okafor 1/5/26): holiday review freeze 12/22 - 1/2, nothing was reviewed in that window. "
    "Anything submitted during the freeze was picked up the week of 1/5. "
    "PLEASE select all columns before sorting this sheet!!"
)


def terms_on(day: date) -> Terms:
    if day >= RATE_SHEET_CHANGE:
        return Terms("9.49", "35.99", "24.99", "33.99", "6.250", "6.471")
    return Terms("8.99", "35.99", "24.99", "33.99", "6.125", "6.338")


# ---------------------------------------------------------------------------------------------
# Calendar helpers
# ---------------------------------------------------------------------------------------------


def is_workday(day: date) -> bool:
    return day.weekday() < SATURDAY and day not in HOLIDAYS


def is_review_day(day: date) -> bool:
    return is_workday(day) and not (FREEZE_START <= day <= FREEZE_END)


def add_review_days(start: date, days: int) -> date:
    current = start
    remaining = days
    while remaining > 0 or not is_review_day(current):
        current += timedelta(days=1)
        if is_review_day(current):
            remaining -= 1
    return current


def add_months(day: date, months: int) -> date:
    month_index = day.month - 1 + months
    year = day.year + month_index // 12
    month = month_index % 12 + 1
    for candidate_day in (day.day, 30, 29, 28):
        try:
            return date(year, month, candidate_day)
        except ValueError:
            continue
    raise ValueError(day)


def months_since_start(day: date) -> int:
    return (day.year - FIRST_SUBMISSION.year) * 12 + day.month - FIRST_SUBMISSION.month


def short_md(day: date) -> str:
    return f"{day.month}/{day.day}"


def short_mdy(day: date) -> str:
    return f"{day.month}/{day.day}/{day.year % 100:02d}"


def no_year(day: date) -> str:
    return f"{MONTH_ABBR[day.month - 1]} {day.day}"


def slug(text: str) -> str:
    return "".join(ch for ch in text.title() if ch.isalnum())


# ---------------------------------------------------------------------------------------------
# Canonical records (the truth the importer should recover)
# ---------------------------------------------------------------------------------------------


@dataclass
class ReviewRound:
    n: int
    on: date
    reviewer: Reviewer
    text: str
    outcome: Outcome


@dataclass
class Placement:
    location: str
    url: str | None
    friendly_text: str | None
    screenshot_url: str | None
    active: bool


@dataclass
class Brief:
    """What the marketer submitted, before review."""

    request_id: str
    submitted: date
    marketer: Marketer
    product: Product
    channel: Channel
    partner: Partner | None
    theme: Theme
    risk: Risk
    issues: list[Issue]
    re_review_of: str | None


@dataclass
class Item:
    brief: Brief
    asset_name: str
    file_name: str
    file_url: str
    trigger_terms: bool
    reviewer: Reviewer | None
    version: int
    round: int
    status: ItemStatus
    outcome: Outcome | None
    rounds: list[ReviewRound]
    newline_comments: bool
    pre_revised: bool
    approval_date: date | None = None
    approver: Reviewer | None = None
    expires: date | None = None
    placement: Placement | None = None

    @property
    def request_id(self) -> str:
        return self.brief.request_id

    @property
    def comments(self) -> str | None:
        if not self.rounds:
            return None
        separator = "\n" if self.newline_comments else " "
        return separator.join(f"R{r.n} {short_md(r.on)} {r.reviewer.initials}: {r.text}" for r in self.rounds)


# ---------------------------------------------------------------------------------------------
# Generation
# ---------------------------------------------------------------------------------------------


@dataclass
class Generator:
    rng: random.Random
    items: list[Item] = field(default_factory=list)
    year_counters: Counter[int] = field(default_factory=Counter)
    re_reviewed: set[str] = field(default_factory=set)

    # -- schedule ---------------------------------------------------------------------------

    def submission_dates(self) -> list[date]:
        dates: list[date] = []
        for year, month, count in MONTHLY_VOLUME:
            first = date(year, month, 1)
            days = [
                first + timedelta(days=i)
                for i in range(31)
                if (first + timedelta(days=i)).month == month
                and is_workday(first + timedelta(days=i))
                and FIRST_SUBMISSION <= first + timedelta(days=i) <= LAST_SUBMISSION
            ]
            picked = sorted(self.rng.choices(days, k=count))
            if FIRST_SUBMISSION in days:
                picked[0] = FIRST_SUBMISSION
            if LAST_SUBMISSION in days:
                picked[-1] = LAST_SUBMISSION
            dates.extend(sorted(picked))
        return dates

    def next_request_id(self, submitted: date, *, reuse_previous: bool) -> str:
        year = submitted.year % 100
        if reuse_previous and self.items:
            return self.items[-1].request_id
        self.year_counters[year] += 1
        return f"MCR-{year:02d}-{self.year_counters[year]:03d}"

    # -- brief ------------------------------------------------------------------------------

    def pick_marketer(self, submitted: date) -> Marketer:
        progress = months_since_start(submitted) / (len(MONTHLY_VOLUME) - 1)
        weights = [m.weight_start + (m.weight_end - m.weight_start) * progress for m in MARKETERS]
        return self.rng.choices(MARKETERS, weights=weights)[0]

    def pick_theme(self, product: Product, submitted: date) -> Theme:
        candidates = [t for t in THEMES if t.product == product and (not t.months or submitted.month in t.months)]
        low = [t for t in candidates if t.low_risk]
        offer = [t for t in candidates if not t.low_risk]
        if low and self.rng.random() < LOW_RISK_THEME_P:
            return self.rng.choice(low)
        return self.rng.choice(offer)

    def pick_partner_slot(self, submitted: date) -> tuple[Partner, Product, Channel]:
        active = [p for p in PARTNERS if p.agreement_date <= submitted]
        partner = self.rng.choices(active, weights=[p.weight for p in active])[0]
        return partner, self.rng.choice(partner.products), self.rng.choice(partner.channels)

    def risk_for(self, product: Product, channel: Channel, partner: Partner | None, theme: Theme) -> Risk:
        # Mirrors the app's intake derivation (spec 4.2.5): partner channels, mortgage and prescreen mail are
        # high; paid ads and landing pages are medium; the rest are low. No-offer brand assets are always low.
        if theme.low_risk:
            return Risk.LOW
        if partner is not None or product == Product.MTG or channel == Channel.DIRECT_MAIL:
            return Risk.HIGH
        if channel in MEDIUM_RISK_CHANNELS:
            return Risk.MEDIUM
        return Risk.LOW

    def eligible_issues(self, brief: Brief) -> list[Issue]:
        out: list[Issue] = []
        for issue in ISSUES:
            if brief.product not in issue.products:
                continue
            if issue.channels and brief.channel not in issue.channels:
                continue
            if issue.partner_only and brief.partner is None:
                continue
            if issue.generic != brief.theme.low_risk:
                continue
            if issue.on_or_after and brief.submitted < issue.on_or_after:
                continue
            if issue.before and brief.submitted >= issue.before:
                continue
            out.append(issue)
        return out

    def initial_issue_count(self, brief: Brief) -> int:
        if brief.theme.low_risk:
            return self.rng.choice((0, 0, 1))
        choices = {Risk.LOW: (0, 1, 1, 2), Risk.MEDIUM: (0, 1, 1, 2, 2), Risk.HIGH: (1, 2, 2, 3, 3)}
        return self.rng.choice(choices[brief.risk])

    def find_re_review_origin(self, submitted: date) -> Item | None:
        start, end = RE_REVIEW_WINDOW
        if not (start <= submitted <= end) or self.rng.random() >= RE_REVIEW_P:
            return None
        horizon = submitted + timedelta(days=RE_REVIEW_LOOKAHEAD_DAYS)
        for item in self.items:
            if item.request_id in self.re_reviewed or item.status != ItemStatus.APPROVED:
                continue
            if item.brief.channel not in WEB_CHANNELS or item.brief.re_review_of is not None:
                continue
            expiring = item.expires is not None and item.expires <= horizon
            stale_rate = item.brief.product == Product.PL and item.brief.submitted < RATE_SHEET_CHANGE <= submitted
            if expiring or stale_rate:
                return item
        return None

    def make_brief(self, submitted: date, *, reuse_id: bool) -> Brief:
        request_id = self.next_request_id(submitted, reuse_previous=reuse_id)
        origin = self.find_re_review_origin(submitted)
        if origin is not None:
            self.re_reviewed.add(origin.request_id)
            src = origin.brief
            brief = Brief(
                request_id, submitted, src.marketer, src.product, src.channel, src.partner, src.theme, src.risk, [],
                origin.request_id,
            )  # fmt: skip
        else:
            marketer = self.pick_marketer(submitted)
            partner: Partner | None = None
            if marketer.partner_led:
                partner, product, channel = self.pick_partner_slot(submitted)
            else:
                product = self.rng.choice(marketer.products)
                channel = self.rng.choice(marketer.channels)
            theme = self.pick_theme(product, submitted)
            risk = self.risk_for(product, channel, partner, theme)
            brief = Brief(request_id, submitted, marketer, product, channel, partner, theme, risk, [], None)
        pool = self.eligible_issues(brief)
        brief.issues = self.rng.sample(pool, min(len(pool), self.initial_issue_count(brief)))
        if origin is not None and brief.product == Product.PL and submitted >= RATE_SHEET_CHANGE:
            stale = next((i for i in pool if i.key in {"stale_floor", "partner_old_rates"}), None)
            if stale is not None and stale not in brief.issues:
                brief.issues.insert(0, stale)
        return brief

    # -- review simulation -------------------------------------------------------------------

    def pick_reviewer(self, brief: Brief, *, exclude: Reviewer | None = None) -> Reviewer:
        active = [r for r in REVIEWERS if r.active_from <= brief.submitted and r != exclude]
        weights = [r.weight * (2.0 if r.is_lead and brief.risk == Risk.HIGH else 1.0) for r in active]
        return self.rng.choices(active, weights=weights)[0]

    def first_wait(self, brief: Brief) -> int:
        base = FIRST_WAIT_BASE + FIRST_WAIT_GROWTH_PER_MONTH * months_since_start(brief.submitted)
        extra = HIGH_RISK_EXTRA_WAIT if brief.risk == Risk.HIGH else 0
        return int(base) + self.rng.randint(0, FIRST_WAIT_JITTER) + extra

    def later_wait(self, brief: Brief, open_issues: Sequence[Issue]) -> int:
        base = LATER_WAIT_BASE + LATER_WAIT_GROWTH_PER_MONTH * months_since_start(brief.submitted)
        extra = SPONSOR_BANK_EXTRA_WAIT if any(i.sponsor_bank for i in open_issues) else 0
        return int(base) + self.rng.randint(0, LATER_WAIT_JITTER) + extra

    def marketer_turnaround(self) -> int:
        low, high = MARKETER_SLOW if self.rng.random() < MARKETER_SLOW_P else MARKETER_FAST
        return self.rng.randint(low, high)

    def flag_text(self, issue: Issue, day: date) -> str:
        terms = terms_on(day)
        return issue.flag.format(
            pl_floor=terms.pl_floor,
            pl_ceil=terms.pl_ceil,
            cc_low=terms.cc_low,
            cc_high=terms.cc_high,
            mtg_rate=terms.mtg_rate,
            mtg_apr=terms.mtg_apr,
        )

    def first_round(self, brief: Brief, open_issues: list[Issue], day: date) -> tuple[Outcome, str]:
        if not open_issues:
            return Outcome.APPROVED, self.rng.choice(APPROVE_PHRASES)
        blocking = [i for i in open_issues if i.blocking]
        if blocking and self.rng.random() < REJECT_P:
            reason = self.flag_text(blocking[0], day)
            return Outcome.REJECTED, f"rejected - {reason}. Can't run this; talk to me before resubmitting."
        if len(open_issues) == 1 and open_issues[0].minor and self.rng.random() < REVIEWER_FIXES_P:
            fixed = open_issues[0]
            open_issues.clear()
            return (
                Outcome.APPROVED_WITH_EDITS,
                f"{self.flag_text(fixed, day)} - fixed it myself, approved w/ changes.",
            )
        text = "; ".join(self.flag_text(i, day) for i in open_issues) + "."
        if brief.partner is not None and self.rng.random() < REVIEWER_FIXES_P:
            text += f" Jordan pls forward to {brief.partner.name}."
        return Outcome.CHANGES_REQUESTED, text

    def later_round(self, brief: Brief, n: int, open_issues: list[Issue], day: date) -> tuple[Outcome, str]:
        resolved = [i for i in open_issues if self.rng.random() < RESOLVE_P or n >= MAX_ROUNDS]
        if not resolved:
            resolved = [open_issues[0]]
        remaining = [i for i in open_issues if i not in resolved]
        ok_part = ", ".join(i.label for i in resolved) + " ok"
        open_issues[:] = remaining
        if not remaining:
            return Outcome.APPROVED, f"{ok_part} - {self.rng.choice(APPROVE_PHRASES)}"
        if len(remaining) == 1 and remaining[0].minor and self.rng.random() < REVIEWER_FIXES_P:
            open_issues.clear()
            return Outcome.APPROVED_WITH_EDITS, (
                f"{ok_part}; {remaining[0].label} still off - fixed it myself on v{n}, approved w/ changes."
            )
        parts = [ok_part, *(i.still for i in remaining)]
        if n <= NEW_ISSUE_MAX_ROUND and self.rng.random() < NEW_ISSUE_P:
            fresh = [i for i in self.eligible_issues(brief) if i not in brief.issues]
            if fresh:
                new_issue = self.rng.choice(fresh)
                brief.issues.append(new_issue)
                open_issues.append(new_issue)
                parts.append(f"new: {self.flag_text(new_issue, day)}")
        return Outcome.CHANGES_REQUESTED, "; ".join(parts) + "."

    def simulate(self, brief: Brief, file_stem: str) -> Item:
        reviewer = self.pick_reviewer(brief)
        assigned_on = add_review_days(brief.submitted, self.rng.randint(0, ASSIGN_DELAY_MAX))
        pre_revised = self.rng.random() < PRE_REVISED_P
        version = 2 if pre_revised else 1
        open_issues = list(brief.issues)
        rounds: list[ReviewRound] = []
        review_day = add_review_days(assigned_on, self.first_wait(brief))
        n = 1
        waiting_on_marketer = False
        while review_day <= AS_OF:
            if n > 1 and self.rng.random() < HANDOFF_P:
                reviewer = self.pick_reviewer(brief, exclude=reviewer)
            if n == 1:
                outcome, text = self.first_round(brief, open_issues, review_day)
            else:
                outcome, text = self.later_round(brief, n, open_issues, review_day)
            if n == 1 and pre_revised:
                text = f"(reviewed v2 - {brief.marketer.initials} sent an updated file before I started) {text}"
            if n == 1 and brief.re_review_of is not None:
                text = f"re-review of {brief.re_review_of}: {text}"
            rounds.append(ReviewRound(n, review_day, reviewer, text, outcome))
            if outcome != Outcome.CHANGES_REQUESTED:
                break
            waiting_on_marketer = True
            too_old = (AS_OF - brief.submitted).days > ABANDON_MIN_AGE_DAYS
            if too_old and self.rng.random() < ABANDON_P:
                break
            resubmitted = add_review_days(review_day, self.marketer_turnaround())
            if resubmitted > AS_OF:
                break
            waiting_on_marketer = False
            n += 1
            version += 1
            review_day = add_review_days(resubmitted, self.later_wait(brief, open_issues))
        return self.finish_item(
            brief,
            file_stem,
            ReviewState(reviewer, assigned_on, version, n, rounds, waiting_on_marketer, pre_revised),
        )

    def finish_item(self, brief: Brief, file_stem: str, state: "ReviewState") -> Item:
        last = state.rounds[-1] if state.rounds else None
        reviewer: Reviewer | None = state.reviewer
        if last is not None and last.outcome in {Outcome.APPROVED, Outcome.APPROVED_WITH_EDITS}:
            status, outcome = ItemStatus.APPROVED, last.outcome
        elif last is not None and last.outcome == Outcome.REJECTED:
            status, outcome = ItemStatus.REJECTED, Outcome.REJECTED
        elif last is not None and state.waiting_on_marketer:
            status, outcome = ItemStatus.CHANGES_REQUESTED, None
        elif last is not None:
            status, outcome = ItemStatus.IN_REVIEW, None
        elif (AS_OF - brief.submitted).days <= UNASSIGNED_WINDOW_DAYS and self.rng.random() < UNASSIGNED_P:
            status, outcome, reviewer = ItemStatus.SUBMITTED, None, None
        else:
            status, outcome = ItemStatus.IN_REVIEW, None
        if last is not None:
            reviewer = last.reviewer if status != ItemStatus.IN_REVIEW else state.reviewer
        trigger = any(i.trigger for i in brief.issues) or (
            not brief.theme.low_risk and self.rng.random() < TRIGGER_WITHOUT_FINDING_P
        )
        ext = CHANNEL_CODE[brief.channel][1]
        file_name = f"{file_stem}_v{state.version}.{ext}"
        item = Item(
            brief=brief,
            asset_name=self.asset_name(brief),
            file_name=file_name,
            file_url=f"{SHAREPOINT}/Submissions/{brief.request_id}/{file_name}",
            trigger_terms=trigger,
            reviewer=reviewer,
            version=state.version,
            round=state.round_n,
            status=status,
            outcome=outcome,
            rounds=state.rounds,
            newline_comments=self.rng.random() < NAME_ONLY_P + 0.2,
            pre_revised=state.pre_revised,
        )
        if status == ItemStatus.APPROVED and last is not None:
            item.approval_date = last.on
            item.approver = last.reviewer
            item.expires = self.expiry_for(brief.channel, last.on)
            item.placement = self.placement_for(item, last.on)
        return item

    def asset_name(self, brief: Brief) -> str:
        noun = self.rng.choice(CHANNEL_NOUNS[brief.channel])
        if brief.re_review_of is not None:
            origin = next(i for i in self.items if i.request_id == brief.re_review_of)
            reason = "rate sheet 6/1" if brief.submitted >= RATE_SHEET_CHANGE else "approval expiring"
            return f"RE-REVIEW ({reason}) {origin.asset_name} - orig {origin.request_id}"
        if brief.partner is not None:
            return f"{brief.partner.name} - {brief.product.code} {brief.theme.name} {noun}"
        return f"{brief.product.code} {brief.theme.name} {noun}"

    def file_stem(self, brief: Brief) -> str:
        code = CHANNEL_CODE[brief.channel][0]
        prefix = slug(brief.partner.name) + "_" if brief.partner is not None else ""
        suffix = "_RR" if brief.re_review_of is not None else ""
        return f"{prefix}{brief.product.code}_{slug(brief.theme.name)}_{code}{suffix}"

    def expiry_for(self, channel: Channel, approved: date) -> date:
        if channel in ONE_OFF_CHANNELS:
            return approved + timedelta(days=self.rng.randint(*SEND_WINDOW_DAYS))
        if channel in EVERGREEN_CHANNELS:
            return add_months(approved, MONTHS_EVERGREEN)
        return add_months(approved, MONTHS_PARTNER)

    def placement_for(self, item: Item, approved: date) -> Placement | None:
        launch = approved + timedelta(days=self.rng.randint(*LAUNCH_DELAY))
        if launch > AS_OF or self.rng.random() < PLACEMENT_UNRECORDED_P:
            return None
        brief = item.brief
        shot = None
        if self.rng.random() < SCREENSHOT_P:
            shot = f"{SHAREPOINT}/Screenshots/{brief.request_id}_live_{launch.isoformat()}.png"
        quarter = f"{launch.year % 100}Q{(launch.month - 1) // 3 + 1}"
        code = f"{brief.product.code}_{slug(brief.theme.name)}_{quarter}"
        channel = brief.channel
        if channel == Channel.AFFILIATE_PAGE and brief.partner is not None:
            url = f"https://www.{brief.partner.domain}/{brief.partner.path}"
            return Placement(url, url, f"{brief.partner.name} listing", shot, active=True)
        if channel == Channel.LANDING_PAGE:
            host = f"https://www.{brief.partner.domain}" if brief.partner is not None else CLEARPATH_SITE
            path = brief.partner.path if brief.partner is not None else f"{brief.product.code.lower()}/{code.lower()}"
            url = f"{host}/{path}"
            return Placement(url, url, "live LP", shot, active=True)
        location = self.offline_location(item, launch, code)
        return Placement(location, None, None, shot, active=channel not in ONE_OFF_CHANNELS)

    def offline_location(self, item: Item, launch: date, code: str) -> str:
        brief = item.brief
        segment = self.rng.choice(SEGMENTS[brief.product])
        if brief.partner is not None and brief.channel == Channel.PAID_SOCIAL:
            return f"IG post @moneywithmaya {short_md(launch)}"
        if brief.partner is not None:
            return f"{brief.partner.name} send {short_md(launch)} (their list)"
        by_channel = {
            Channel.EMAIL: f"Sent {short_md(launch)} to {segment}",
            Channel.SMS: f"Sent {short_md(launch)} - {segment}",
            Channel.DIRECT_MAIL: f"Drop {short_md(launch)}, {self.rng.randint(12, 90)}k pieces",
            Channel.PAID_SOCIAL: f"Meta campaign {code}",
            Channel.PAID_SEARCH: f"Google Ads - {code}",
            Channel.DISPLAY: f"Programmatic line item {code}",
            Channel.CALL_SCRIPT: f"Live in call center KB-{self.rng.randint(1100, 1999)}",
        }
        return by_channel[brief.channel]

    def build(self) -> list[Item]:
        for index, submitted in enumerate(self.submission_dates()):
            brief = self.make_brief(submitted, reuse_id=index == REUSED_ID_INDEX)
            self.items.append(self.simulate(brief, self.file_stem(brief)))
        return self.items


@dataclass
class ReviewState:
    reviewer: Reviewer
    assigned_on: date
    version: int
    round_n: int
    rounds: list[ReviewRound]
    waiting_on_marketer: bool
    pre_revised: bool


# ---------------------------------------------------------------------------------------------
# Rendering canonical items into spreadsheet rows
# ---------------------------------------------------------------------------------------------

type CellScalar = str | int | date | None


@dataclass
class CellValue:
    value: CellScalar
    hyperlink: str | None = None


class RowKind(StrEnum):
    ITEM = "item"
    DUPLICATE = "duplicate"
    NOTE = "note"


@dataclass
class SheetRow:
    kind: RowKind
    item: Item | None
    cells: list[CellValue]
    excel_row: int = 0


def render_asset(item: Item, rng: random.Random) -> CellValue:
    early = item.brief.submitted < HYPERLINK_ERA
    if rng.random() < (HYPERLINK_P_EARLY if early else HYPERLINK_P_LATE):
        return CellValue(item.file_name, item.file_url)
    if rng.random() < NAME_ONLY_P:
        return CellValue(f"{item.asset_name} (emailed {short_md(item.brief.submitted)})")
    return CellValue(f"{item.asset_name} - {item.file_url}")


def render_live(item: Item, rng: random.Random) -> CellValue:
    placement = item.placement
    if placement is None:
        return CellValue(None)
    early = item.brief.submitted < HYPERLINK_ERA
    if placement.url is not None and rng.random() < (HYPERLINK_P_EARLY if early else HYPERLINK_P_LATE):
        friendly = placement.friendly_text if rng.random() < FRIENDLY_LINK_TEXT_P else None
        return CellValue(friendly or placement.url, placement.url)
    if placement.screenshot_url is not None:
        return CellValue(f"{placement.location}; screenshot: {placement.screenshot_url}")
    return CellValue(placement.location)


def render_item(item: Item, rng: random.Random) -> list[CellValue]:
    brief = item.brief
    approval = None
    if item.approval_date is not None and item.approver is not None:
        approval = f"{short_mdy(item.approval_date)} - {item.approver.approver_label}"
    status_label = STATUS_DEFAULT_LABEL[(item.status, item.outcome)]
    values: list[CellValue] = [
        CellValue(brief.request_id),
        CellValue(brief.submitted),
        CellValue(brief.marketer.submitter_cell),
        CellValue(PRODUCT_LABEL[brief.product]),
        CellValue(str(brief.channel)),
        render_asset(item, rng),
        CellValue(brief.partner.name if brief.partner is not None else None),
        CellValue(str(brief.risk)),
        CellValue("Y" if item.trigger_terms else "N"),
        CellValue(item.reviewer.name if item.reviewer is not None else None),
        CellValue(item.version),
        CellValue(item.round),
        CellValue(status_label),
        CellValue(item.comments),
        CellValue(approval),
        CellValue(item.expires),
        render_live(item, rng),
    ]
    return values


def layout_rows(items: list[Item], rng: random.Random) -> list[SheetRow]:
    rows: list[SheetRow] = []
    first_2026 = next(i for i, item in enumerate(items) if item.brief.submitted.year == FREEZE_END.year)
    for index, item in enumerate(items):
        if index == first_2026:
            note_cells = [CellValue(NOTE_ROW_TEXT)] + [CellValue(None) for _ in HEADERS[1:]]
            rows.append(SheetRow(RowKind.NOTE, None, note_cells))
        rows.append(SheetRow(RowKind.ITEM, item, render_item(item, rng)))
        if index == DUPLICATE_SOURCE_INDEX + DUPLICATE_OFFSET:
            source = next(r for r in rows if r.item is items[DUPLICATE_SOURCE_INDEX])
            rows.append(SheetRow(RowKind.DUPLICATE, source.item, source.cells))
    for position, row in enumerate(rows):
        row.excel_row = position + 2
    return rows


# ---------------------------------------------------------------------------------------------
# Intended messiness. Every change is recorded; the README oracle is generated from the records.
# ---------------------------------------------------------------------------------------------


class MessKind(StrEnum):
    STATUS = "Status spelling variants"
    PRODUCT = "Product spelling variants"
    DATE_SUBMITTED = "Date Submitted stored as text"
    EXPIRY_TEXT = "Expiration stored as text"
    APPROVER_ONLY = "Approval cell with approver only (no date)"
    BLANK_EXPIRY = "Approved row with blank expiration"
    DUPLICATE_ROW = "Exact duplicate row"
    REUSED_ID = "Request ID reused for a different asset"
    WHITESPACE = "Leading / trailing whitespace"
    PARTNER_NAME = "Partner typed differently from the Partners sheet"
    TRIGGER_YN = "Trigger Terms Y/N variants"
    NOTE_ROW = "Merged note row inside the data"


@dataclass(frozen=True)
class MessRecord:
    kind: MessKind
    request_id: str
    excel_rows: tuple[int, ...]
    column: str
    raw: str
    expected: str


@dataclass
class MessPlan:
    rng: random.Random
    rows: list[SheetRow]
    records: list[MessRecord] = field(default_factory=list)
    touched: set[tuple[int, Col]] = field(default_factory=set)
    protected: set[int] = field(default_factory=set)

    def candidates(self, column: Col, predicate: Callable[[SheetRow], bool]) -> list[SheetRow]:
        return [
            row
            for row in self.rows
            if row.kind == RowKind.ITEM
            and row.excel_row not in self.protected
            and (row.excel_row, column) not in self.touched
            and predicate(row)
        ]

    def mutate(self, plan: "Mutation") -> None:
        pool = self.candidates(plan.column, plan.predicate)
        for row in self.rng.sample(pool, min(plan.count, len(pool))):
            item = row.item
            if item is None:
                continue
            new_value, expected = plan.transform(row)
            row.cells[plan.column] = CellValue(new_value, row.cells[plan.column].hyperlink)
            self.touched.add((row.excel_row, plan.column))
            self.records.append(
                MessRecord(
                    plan.kind, item.request_id, (row.excel_row,), HEADERS[plan.column], display(new_value), expected
                )
            )


@dataclass(frozen=True)
class Mutation:
    kind: MessKind
    column: Col
    count: int
    predicate: Callable[[SheetRow], bool]
    transform: Callable[[SheetRow], tuple[CellScalar, str]]


def display(value: CellScalar) -> str:
    if value is None:
        return "(empty)"
    if isinstance(value, date):
        return f"date {value.isoformat()}"
    return '"' + str(value).replace("\n", "\\n") + '"'


def item_of(row: SheetRow) -> Item:
    if row.item is None:
        raise ValueError(row)
    return row.item


def status_is(status: ItemStatus, outcome: Outcome | None = None) -> Callable[[SheetRow], bool]:
    def check(row: SheetRow) -> bool:
        item = item_of(row)
        return item.status == status and (outcome is None or item.outcome == outcome)

    return check


def product_is(product: Product) -> Callable[[SheetRow], bool]:
    return lambda row: item_of(row).brief.product == product


def canonical_status(item: Item) -> str:
    if item.outcome in {Outcome.APPROVED, Outcome.APPROVED_WITH_EDITS, Outcome.REJECTED}:
        return f"{item.status} (decision: {item.outcome})"
    return str(item.status)


def constant(value: CellScalar, expected: Callable[[Item], str]) -> Callable[[SheetRow], tuple[CellScalar, str]]:
    return lambda row: (value, expected(item_of(row)))


def status_mutations() -> list[Mutation]:
    approved = status_is(ItemStatus.APPROVED, Outcome.APPROVED)
    awc = status_is(ItemStatus.APPROVED, Outcome.APPROVED_WITH_EDITS)
    changes = status_is(ItemStatus.CHANGES_REQUESTED)
    in_review = status_is(ItemStatus.IN_REVIEW)
    submitted = status_is(ItemStatus.SUBMITTED)
    n = STATUS_VARIANTS_PER_KIND
    plan = [
        (APPROVED_TRAILING_SPACE, approved, "approved "),
        (LOWERCASE_AWC_COUNT, awc, "Approved w/ changes"),
        (MINOR_EDITS_COUNT, awc, "Approved - minor edits"),
        (n, changes, "Changes Req."),
        (n, changes, "Needs changes"),
        (n, in_review, "in review"),
        (n, submitted, "Pending"),
        (n, submitted, None),
    ]
    return [
        Mutation(MessKind.STATUS, Col.STATUS, count, predicate, constant(raw, canonical_status))
        for count, predicate, raw in plan
    ]


def product_mutations() -> list[Mutation]:
    plan = [
        (PRODUCT_ABBREV_PL, Product.PL, "PL"),
        (PRODUCT_LOWER_PL, Product.PL, "personal loans"),
        (PRODUCT_ABBREV_CC, Product.CC, "CC"),
        (PRODUCT_ABBREV_MTG, Product.MTG, "MTG"),
    ]
    return [
        Mutation(MessKind.PRODUCT, Col.PRODUCT, count, product_is(product), constant(raw, lambda i: i.brief.product))
        for count, product, raw in plan
    ]


def date_mutations() -> list[Mutation]:
    def as_text(fmt: Callable[[date], str]) -> Callable[[SheetRow], tuple[CellScalar, str]]:
        return lambda row: (fmt(item_of(row).brief.submitted), item_of(row).brief.submitted.isoformat())

    def no_year_text(row: SheetRow) -> tuple[CellScalar, str]:
        submitted = item_of(row).brief.submitted
        return no_year(submitted), f"{submitted.isoformat()} (year from the MCR-YY- prefix)"

    def expiry_text(row: SheetRow) -> tuple[CellScalar, str]:
        expires = item_of(row).expires
        if expires is None:
            raise ValueError(row)
        return short_mdy(expires), expires.isoformat()

    def any_item(_row: SheetRow) -> bool:
        return True

    def has_expiry(row: SheetRow) -> bool:
        return item_of(row).expires is not None

    return [
        Mutation(MessKind.DATE_SUBMITTED, Col.DATE_SUBMITTED, DATE_SLASH_TEXT, any_item, as_text(short_mdy)),
        Mutation(MessKind.DATE_SUBMITTED, Col.DATE_SUBMITTED, DATE_ISO_TEXT, any_item, as_text(date.isoformat)),
        Mutation(MessKind.DATE_SUBMITTED, Col.DATE_SUBMITTED, DATE_NO_YEAR, any_item, no_year_text),
        Mutation(MessKind.EXPIRY_TEXT, Col.EXPIRATION, EXPIRY_TEXT, has_expiry, expiry_text),
    ]


def approval_mutations() -> list[Mutation]:
    def approved_with_expiry(row: SheetRow) -> bool:
        return item_of(row).status == ItemStatus.APPROVED and item_of(row).expires is not None

    def approver_only(row: SheetRow) -> tuple[CellScalar, str]:
        item = item_of(row)
        if item.approver is None or item.approval_date is None:
            raise ValueError(row)
        expected = (
            f"approver {item.approver.name}; date {item.approval_date.isoformat()} from the last R-segment in Comments"
        )
        return item.approver.approver_label, expected

    def blank_expiry(row: SheetRow) -> tuple[CellScalar, str]:
        return None, f"expires_at null + 'missing expiry' warning (policy default would be {item_of(row).expires})"

    return [
        Mutation(MessKind.APPROVER_ONLY, Col.APPROVAL, APPROVER_ONLY, approved_with_expiry, approver_only),
        Mutation(MessKind.BLANK_EXPIRY, Col.EXPIRATION, BLANK_EXPIRY, approved_with_expiry, blank_expiry),
    ]


def yn_mutations() -> list[Mutation]:
    def truth(item: Item) -> str:
        return f"null (unknown) + warning; generator truth was {'Y' if item.trigger_terms else 'N'}"

    def is_yes(row: SheetRow) -> bool:
        return item_of(row).trigger_terms

    def any_item(_row: SheetRow) -> bool:
        return True

    return [
        Mutation(MessKind.TRIGGER_YN, Col.TRIGGER, YN_YES, is_yes, constant("yes", lambda _i: "true")),
        Mutation(MessKind.TRIGGER_YN, Col.TRIGGER, YN_BLANK, any_item, constant(None, truth)),
        Mutation(MessKind.TRIGGER_YN, Col.TRIGGER, YN_UNKNOWN, any_item, constant("?", truth)),
    ]


def partner_mutations() -> list[Mutation]:
    def partner_named(name: str) -> Callable[[SheetRow], bool]:
        def check(row: SheetRow) -> bool:
            partner = item_of(row).brief.partner
            return partner is not None and partner.name == name

        return check

    def canonical_partner(item: Item) -> str:
        return f"partner '{item.brief.partner.name if item.brief.partner else ''}'"

    plan = [
        (2, "LendCompare", "Lend Compare"),
        (1, "CardPerks Weekly", "CardPerks"),
        (1, "LeadStream Direct", "Leadstream Direct"),
    ]
    return [
        Mutation(MessKind.PARTNER_NAME, Col.PARTNER, count, partner_named(name), constant(raw, canonical_partner))
        for count, name, raw in plan
    ]


def whitespace_mutations() -> list[Mutation]:
    def pad(column: Col, template: str) -> Callable[[SheetRow], tuple[CellScalar, str]]:
        def transform(row: SheetRow) -> tuple[CellScalar, str]:
            value = str(row.cells[column].value)
            return template.format(value), f'"{value}" after strip()'

        return transform

    def has_partner(row: SheetRow) -> bool:
        return item_of(row).brief.partner is not None

    def has_reviewer(row: SheetRow) -> bool:
        return item_of(row).reviewer is not None

    def any_item(_row: SheetRow) -> bool:
        return True

    return [
        Mutation(MessKind.WHITESPACE, Col.REQUEST_ID, 1, any_item, pad(Col.REQUEST_ID, "{} ")),
        Mutation(MessKind.WHITESPACE, Col.SUBMITTER, 2, any_item, pad(Col.SUBMITTER, "{}  ")),
        Mutation(MessKind.WHITESPACE, Col.PRODUCT, 1, any_item, pad(Col.PRODUCT, " {}")),
        Mutation(MessKind.WHITESPACE, Col.CHANNEL, 1, any_item, pad(Col.CHANNEL, "{} ")),
        Mutation(MessKind.WHITESPACE, Col.PARTNER, 2, has_partner, pad(Col.PARTNER, " {} ")),
        Mutation(MessKind.WHITESPACE, Col.REVIEWER, 2, has_reviewer, pad(Col.REVIEWER, "{} ")),
    ]


def blank_status_with_approval(plan: MessPlan) -> None:
    pool = plan.candidates(Col.STATUS, status_is(ItemStatus.APPROVED, Outcome.APPROVED))
    pool = [row for row in pool if (row.excel_row, Col.APPROVAL) not in plan.touched]
    row = plan.rng.choice(pool)
    item = item_of(row)
    row.cells[Col.STATUS] = CellValue(None)
    plan.touched.add((row.excel_row, Col.STATUS))
    plan.records.append(
        MessRecord(
            MessKind.STATUS,
            item.request_id,
            (row.excel_row,),
            HEADERS[Col.STATUS],
            display(None),
            f"{canonical_status(item)} - inferred from the filled Approval cell; emit a warning",
        )
    )


def apply_mess(rows: list[SheetRow], items: list[Item], rng: random.Random) -> list[MessRecord]:
    plan = MessPlan(rng, rows)
    dup_rows = [r for r in rows if r.item is items[DUPLICATE_SOURCE_INDEX]]
    reused_rows = [r for r in rows if r.item is items[REUSED_ID_INDEX] or r.item is items[REUSED_ID_INDEX - 1]]
    plan.protected = {r.excel_row for r in (*dup_rows, *reused_rows)}
    mutations = [
        *status_mutations(),
        *product_mutations(),
        *date_mutations(),
        *approval_mutations(),
        *yn_mutations(),
        *partner_mutations(),
        *whitespace_mutations(),
    ]
    for mutation in mutations:
        plan.mutate(mutation)
    blank_status_with_approval(plan)
    original = items[DUPLICATE_SOURCE_INDEX]
    plan.records.append(
        MessRecord(
            MessKind.DUPLICATE_ROW,
            original.request_id,
            tuple(r.excel_row for r in dup_rows),
            "(whole row)",
            "every cell identical, including hyperlinks",
            "import one item; report the second row as a duplicate and skip it",
        )
    )
    first, second = items[REUSED_ID_INDEX - 1], items[REUSED_ID_INDEX]
    plan.records.append(
        MessRecord(
            MessKind.REUSED_ID,
            first.request_id,
            tuple(r.excel_row for r in reused_rows),
            HEADERS[Col.REQUEST_ID],
            display(first.request_id),
            f"two items: r{reused_rows[0].excel_row} '{first.asset_name}' (submitted {first.brief.submitted}) and "
            f"r{reused_rows[1].excel_row} '{second.asset_name}' (submitted {second.brief.submitted}). Import both; "
            f"the later row gets a disambiguated legacy_ref (e.g. {first.request_id}-b) and an admin warning",
        )
    )
    note = next(r for r in rows if r.kind == RowKind.NOTE)
    plan.records.append(
        MessRecord(
            MessKind.NOTE_ROW,
            "(none)",
            (note.excel_row,),
            f"A{note.excel_row}:Q{note.excel_row} merged",
            display(NOTE_ROW_TEXT[:60] + "..."),
            "skip: column A does not match the Request ID pattern; report it as a skipped non-data row",
        )
    )
    return plan.records


# ---------------------------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------------------------

FONT_NAME = "Arial"
BODY_FONT = Font(name=FONT_NAME, size=10)
LINK_FONT = Font(name=FONT_NAME, size=10, color="0563C1", underline="single")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(name=FONT_NAME, size=10, bold=True, italic=True, color="7F6000")
THIN = Side(style="thin", color="BFBFBF")
HEADER_BORDER = Border(bottom=THIN)
DATE_FORMAT = "m/d/yyyy"
HEADER_HEIGHT = 32
STATUS_COLORS: tuple[tuple[str, str, str], ...] = (
    # (formula, fill, font colour); first match wins, so "Approved w/ changes" never turns orange
    ('ISNUMBER(SEARCH("approv",$M2))', "C6EFCE", "006100"),
    ('ISNUMBER(SEARCH("reject",$M2))', "FFC7CE", "9C0006"),
    ('ISNUMBER(SEARCH("change",$M2))', "FCE4D6", "843C0C"),
    ('ISNUMBER(SEARCH("review",$M2))', "FFEB9C", "9C5700"),
    ('OR(ISNUMBER(SEARCH("pending",$M2)),ISNUMBER(SEARCH("not started",$M2)))', "D9D9D9", "404040"),
)


@dataclass(frozen=True)
class ListColumn:
    header: str
    values: tuple[str, ...]
    validates: Col | None


LIST_COLUMNS: tuple[ListColumn, ...] = (
    ListColumn("Status", STATUS_LIST, Col.STATUS),
    ListColumn("Product", tuple(PRODUCT_LABEL.values()), Col.PRODUCT),
    ListColumn("Channel", tuple(str(c) for c in Channel), Col.CHANNEL),
    ListColumn("Risk Tier", tuple(str(r) for r in Risk), Col.RISK),
    ListColumn("Trigger Terms (Y/N)", ("Y", "N"), Col.TRIGGER),
    ListColumn("Reviewer", tuple(r.name for r in REVIEWERS), Col.REVIEWER),
    ListColumn("Business Unit", tuple(m.bu for m in MARKETERS), None),
)


def style_header(ws: Worksheet, headers: Sequence[str], widths: Sequence[int]) -> None:
    for index, (title, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = ws.cell(row=1, column=index, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[LETTERS[index - 1]].width = width
    ws.row_dimensions[1].height = HEADER_HEIGHT
    ws.freeze_panes = "A2"


def write_tracker(ws: Worksheet, rows: list[SheetRow]) -> None:
    style_header(ws, HEADERS, COLUMN_WIDTHS)
    for row in rows:
        if row.kind == RowKind.NOTE:
            cell = ws.cell(row=row.excel_row, column=1, value=row.cells[0].value)
            cell.font = NOTE_FONT
            cell.fill = NOTE_FILL
            cell.alignment = Alignment(vertical="center")
            ws.merge_cells(f"A{row.excel_row}:Q{row.excel_row}")
            continue
        for column, value in enumerate(row.cells):
            cell = ws.cell(row=row.excel_row, column=column + 1, value=value.value)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=column == Col.COMMENTS)
            if isinstance(value.value, date):
                cell.number_format = DATE_FORMAT
            if value.hyperlink is not None:
                cell.hyperlink = value.hyperlink
                cell.font = LINK_FONT
    last = rows[-1].excel_row
    ws.auto_filter.ref = f"A1:Q{last}"
    status_range = f"M2:M{last}"
    for formula, fill, font in STATUS_COLORS:
        rule = FormulaRule(
            formula=[formula],
            fill=PatternFill("solid", fgColor=fill, bgColor=fill),
            font=Font(color=font),
            stopIfTrue=True,
        )
        ws.conditional_formatting.add(status_range, rule)
    for index, column in enumerate(LIST_COLUMNS):
        if column.validates is None:
            continue
        letter = LETTERS[index]
        # showErrorMessage=False: the dropdown is offered but anything typed or pasted is accepted, which is
        # how a real tracker drifts away from its own lists.
        validation = DataValidation(
            type="list",
            formula1=f"=Lists!${letter}$2:${letter}${len(column.values) + 1}",
            allow_blank=True,
            showErrorMessage=False,
        )
        target = LETTERS[column.validates]
        validation.add(f"{target}2:{target}{last}")
        ws.add_data_validation(validation)


def write_lists(ws: Worksheet) -> None:
    style_header(ws, [c.header for c in LIST_COLUMNS], [22] * len(LIST_COLUMNS))
    for column_index, column in enumerate(LIST_COLUMNS, start=1):
        for row_index, value in enumerate(column.values, start=2):
            ws.cell(row=row_index, column=column_index, value=value).font = BODY_FONT


PARTNER_HEADERS = ("Partner Name", "Type", "Description", "Contact Email", "Agreement Date", "Status", "Notes")


def write_partners(ws: Worksheet) -> None:
    style_header(ws, PARTNER_HEADERS, (22, 13, 52, 36, 15, 10, 60))
    for row_index, partner in enumerate(PARTNERS, start=2):
        values: tuple[CellScalar, ...] = (
            partner.name,
            PARTNER_KIND_LABEL[partner.kind],
            partner.description,
            partner.contact_email,
            partner.agreement_date,
            partner.status,
            partner.notes or None,
        )
        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=value)
            cell.font = BODY_FONT
            if isinstance(value, date):
                cell.number_format = DATE_FORMAT
    ws.auto_filter.ref = f"A1:G{len(PARTNERS) + 1}"


def build_workbook(rows: list[SheetRow]) -> Workbook:
    wb = Workbook()
    tracker = wb.active
    if tracker is None:
        raise RuntimeError("new workbook has no active sheet")
    tracker.title = "Review Tracker"
    write_tracker(tracker, rows)
    write_lists(wb.create_sheet("Lists"))
    write_partners(wb.create_sheet("Partners"))
    wb.properties.creator = "Rachel Okafor"
    wb.properties.lastModifiedBy = "Dana Kim"
    wb.properties.title = "Marketing Compliance Review Tracker"
    wb.properties.created = FILE_CREATED
    # openpyxl stamps `modified` (and zip entry times) with the wall clock at save, so bytes are not stable.
    return wb


# ---------------------------------------------------------------------------------------------
# README / import oracle
# ---------------------------------------------------------------------------------------------

README_INTRO = """\
# Sample legacy review tracker

`clearpath_review_tracker.xlsx` is the fictional Excel tracker ClearPath's compliance team used before the
tool: one row per review request, redlines typed into one cell, approvals buried in email. The Import page
ingests it ("Load sample tracker") to prove migration works, and its history seeds queue and metrics data.

**Generated, not hand-edited.** `uv run scripts/generate_sample_tracker.py` rewrites both this README and
the workbook from one seeded RNG; re-running produces identical cell values. Change the script, then
regenerate. Every row number and Request ID below is computed from the same data the workbook is written
from, so this file is the test oracle for the importer (spec 4.3, Import: "Every intended mess listed in
`samples/README.md` is handled as specified there").

Everything is fictional. Domains use the reserved `.example` TLD; the sponsor bank (Harborview Bank, N.A.),
license and NMLS numbers are invented.
"""

SHEET_ROWS: tuple[tuple[str, str], ...] = (
    (
        "`Review Tracker`",
        (
            "Header in row 1 (frozen, autofiltered, dark fill). One row per request, plus one merged note row and one "
            "pasted duplicate (see below). Status is colour-coded with `SEARCH()`-based conditional formats, so "
            "misspelled statuses still colour. Data validation dropdowns point at `Lists` but have "
            "`showErrorMessage=False`, so off-list values were accepted."
        ),
    ),
    (
        "`Lists`",
        (
            "The dropdown values the team *intended* to use: Status, Product, Channel, Risk Tier, Y/N, Reviewer, "
            "Business Unit. The data does not always follow them."
        ),
    ),
    (
        "`Partners`",
        "Partner name, type, description, contact email, agreement date, status, notes. Names here are canonical.",
    ),
)

COLUMN_ROWS: tuple[tuple[str, str], ...] = (
    (
        "Request ID",
        (
            "`MCR-YY-NNN`, numbered per year in submission order (regex `^MCR-\\d{2}-\\d{3}$` after strip). Becomes "
            "`legacy_ref`; it is not the app's `CP-####` display ID."
        ),
    ),
    (
        "Date Submitted",
        "A real Excel date (`m/d/yyyy` format; openpyxl loads it as a `datetime` at midnight). Exceptions below.",
    ),
    ("Submitter / BU", "`First Last / Business Unit`."),
    ("Product", "A `Lists` label (`Personal Loan`, `Credit Card`, `Mortgage Prequal`); app values below."),
    ("Channel", "One of the nine `Lists` channels."),
    (
        "Asset Name / File Link",
        (
            "One of three shapes: a hyperlinked file name (target = SharePoint URL), "
            "`<asset name> - <SharePoint URL>` as text, or `<asset name> (emailed M/D)` with no link at all."
        ),
    ),
    ("Affiliate / Partner", "Blank for first-party assets, else a Partners-sheet name."),
    (
        "Risk Tier",
        (
            "`High` / `Medium` / `Low`, following the app's intake derivation: partner, mortgage and direct mail "
            "high; paid ads and LPs medium; the rest low; no-offer brand assets low."
        ),
    ),
    ("Trigger Terms Present (Y/N)", "`Y` / `N`: whether the copy contains a Reg Z trigger term."),
    ("Assigned Reviewer", "Full name from `Lists`; blank when nobody has picked the item up."),
    (
        "Version #",
        (
            "Latest version the marketer sent. Equals Revision Round, except where the marketer sent v2 before the "
            "first review (listed below)."
        ),
    ),
    (
        "Revision Round",
        (
            "Current or last review round, starting at 1. Equals the number of comment segments, or that plus 1 when "
            "the item was resubmitted and the new round has not been reviewed yet (status In Review). Unreviewed "
            "items have round 1 and no comments."
        ),
    ),
    ("Status", 'A `Lists` status; mapping in "Status values" below. The outcome comes from here, not from Comments.'),
    (
        "Reviewer Comments / Required Changes",
        "Every round in one cell; grammar below. Blank when no review has happened yet.",
    ),
    (
        "Approval Date & Approver",
        (
            "`M/D/YY - F. Last` on approved rows, blank otherwise. The date equals the last comment segment's date and "
            "the approver is that segment's reviewer."
        ),
    ),
    (
        "Approval Expiration / Re-review Date",
        (
            "Real Excel date. The policy the team followed: one-off sends (Email, SMS, Direct Mail) = end of the send "
            "window, 14-45 days after approval; Landing Page and Call Script = 12 months; Affiliate Page, Paid "
            "Social, Paid Search and Display = 6 months."
        ),
    ),
    (
        "Live URL / Placement & Screenshot Link",
        (
            "Blank, a hyperlink (target = placement URL; display text is the URL or a friendly label), a plain-text "
            "URL, or a placement description (`Sent 3/18 to ...`, `Drop 4/2, 48k pieces`, `Meta campaign ...`, "
            "`IG post @moneywithmaya 2/14`, `Live in call center KB-1234`). Text cells may end in "
            "`; screenshot: <url>`."
        ),
    ),
)

COMMENT_GRAMMAR = """\
### Comment grammar

Each round is one segment: `R<n> <M>/<D> <II>: <text>`. Segments are separated by a single space or by a
newline (Alt+Enter) depending on who typed them; both appear. Parse with

```
R(?P<round>\\d+) (?P<month>\\d{1,2})/(?P<day>\\d{1,2}) (?P<initials>[A-Z]{2}):
```

and split on the match starts; segment text never contains that pattern. Dates have no year: start from
the Date Submitted year and add one each time the month goes backwards (rows that cross New Year are listed
below). Initials are unique: DK Dana Kim, SO Samuel Ortega, LB Leah Brandt, NR Nikhil Rao (joined 3/2/26),
RO Rachel Okafor (compliance lead). On decided rows the last segment is the decision round (its date is the
decision date); every earlier segment requested changes. Do not classify segments by searching for
"approved": findings quote copy such as 'You're pre-approved!'. Take the outcome from Status.
"""


def md_table(headers: tuple[str, str], body: Sequence[tuple[str, str]]) -> str:
    lines = [f"| {headers[0]} | {headers[1]} |", "|---|---|"]
    lines.extend(f"| {left} | {right} |" for left, right in body)
    return "\n".join(lines) + "\n"


README_EXPECTATIONS: dict[MessKind, str] = {
    MessKind.STATUS: "Map case- and whitespace-insensitively through the status table below. A blank Status is "
    "approved (decision approved, with a warning) when the Approval cell is filled, otherwise submitted. Values "
    "outside the table are a preview problem, not a guess.",
    MessKind.PRODUCT: "Strip, then map case-insensitively: PL / Personal Loan / personal loans -> personal_loan, "
    "CC / Credit Card -> credit_card, MTG / Mortgage Prequal -> mortgage_prequal.",
    MessKind.DATE_SUBMITTED: "Accept real dates, `M/D/YY` (US order, 20YY), ISO `YYYY-MM-DD`, and `Mon D` with the "
    "year taken from the Request ID prefix. Report each coerced cell as an info-level note.",
    MessKind.EXPIRY_TEXT: "Same parser as Date Submitted (`M/D/YY` text).",
    MessKind.APPROVER_ONLY: "Approver resolved from `F. Last` to the reviewer; approval date taken from the final "
    "comment segment (and flagged as derived).",
    MessKind.BLANK_EXPIRY: "Store `expires_at` null and report it. Do not invent a date (the generator's policy "
    "default is shown only so tests can tell the rows apart).",
    MessKind.DUPLICATE_ROW: "Rows identical in every cell are one item.",
    MessKind.REUSED_ID: "Same Request ID, different Asset Name / Date Submitted: two items. The discriminator is "
    "(Asset Name, Date Submitted); identical values mean duplicate, different mean reuse.",
    MessKind.WHITESPACE: "Strip every string cell before matching or storing.",
    MessKind.PARTNER_NAME: "Match partners by a normalized key (casefold, drop spaces and punctuation) against the "
    "Partners sheet; the bare `CardPerks` resolves to `CardPerks Weekly` because it is a unique prefix. "
    "Report each fuzzy match.",
    MessKind.TRIGGER_YN: "`Y`/`yes` -> true, `N` -> false, blank or `?` -> null with a warning. Do not infer from "
    "comments.",
    MessKind.NOTE_ROW: "Skip rows whose stripped column A does not match the Request ID pattern; report them.",
}


def md_escape(text: str) -> str:
    return text.replace("|", "\\|")


def refs(rows: Sequence[SheetRow]) -> str:
    return ", ".join(f"{item_of(r).request_id.strip()} (r{r.excel_row})" for r in rows)


def data_rows(rows: Sequence[SheetRow]) -> list[SheetRow]:
    return [r for r in rows if r.kind in {RowKind.ITEM, RowKind.DUPLICATE}]


@dataclass(frozen=True)
class ValueTable:
    title: str
    column: Col
    canonical: Callable[[Item], str]
    clean_values: frozenset[str]


def value_table(rows: Sequence[SheetRow], table: ValueTable) -> list[str]:
    groups: dict[tuple[str, str], list[SheetRow]] = {}
    for row in data_rows(rows):
        raw = display(row.cells[table.column].value)
        groups.setdefault((raw, table.canonical(item_of(row))), []).append(row)
    lines = [
        f"### {table.title}",
        "",
        "| Raw cell | Canonical | Rows | Where (off-list values only) |",
        "|---|---|---|---|",
    ]
    for (raw, canonical), members in sorted(groups.items()):
        where = "" if raw in table.clean_values else refs(members)
        lines.append(f"| `{md_escape(raw)}` | {canonical} | {len(members)} | {where} |")
    lines.append("")
    return lines


def mess_section(records: Sequence[MessRecord]) -> list[str]:
    lines = ["## Intended messiness (the oracle)", ""]
    lines.append(
        "Each kind lists every affected cell. `rN` is the Excel row number. Values in backticks are exact, "
        "including spaces; `(empty)` is an empty cell. 'Importer should' is the required behaviour."
    )
    lines.append("")
    for kind in MessKind:
        kind_records = [r for r in records if r.kind == kind]
        lines.extend([f"### {kind}", "", f"Importer should: {README_EXPECTATIONS[kind]}", ""])
        lines.extend(["| Request ID | Row(s) | Column | Raw | Expected |", "|---|---|---|---|---|"])
        for record in sorted(kind_records, key=lambda r: (r.excel_rows, r.column)):
            rows_text = ", ".join(f"r{n}" for n in record.excel_rows)
            lines.append(
                f"| {record.request_id.strip()} | {rows_text} | {record.column} | `{md_escape(record.raw)}` "
                f"| {md_escape(record.expected)} |"
            )
        lines.append("")
    return lines


def link_section(rows: Sequence[SheetRow]) -> list[str]:
    lines = ["### Links: hyperlink objects vs text", ""]
    lines.append(
        "Importer should: read `cell.hyperlink.target` when a hyperlink exists (the display text may be a label "
        "such as `LendCompare listing`), otherwise extract URLs from the text with a regex. Hyperlinked cells:"
    )
    lines.append("")
    for column in (Col.ASSET, Col.LIVE):
        linked = [r for r in data_rows(rows) if r.cells[column].hyperlink is not None]
        relabeled = [r for r in linked if r.cells[column].value != r.cells[column].hyperlink]
        text_urls = [
            r
            for r in data_rows(rows)
            if r.cells[column].hyperlink is None and "https://" in str(r.cells[column].value or "")
        ]
        lines.append(f"- **{HEADERS[column]}**: {len(linked)} hyperlinked cells: {refs(linked)}.")
        if column == Col.LIVE:
            lines.append(f"  - Display text differs from the target (friendly label): {refs(relabeled)}.")
        lines.append(f"  - Text containing a URL, no hyperlink object ({len(text_urls)}): {refs(text_urls)}.")
        if column == Col.ASSET:
            no_link = [
                r
                for r in data_rows(rows)
                if r.cells[column].hyperlink is None and "https://" not in str(r.cells[column].value or "")
            ]
            lines.append(f"  - Asset name only, no link at all ({len(no_link)}): {refs(no_link)}.")
    lines.append("")
    return lines


def characteristics_section(
    rows: Sequence[SheetRow], items: Sequence[Item], records: Sequence[MessRecord]
) -> list[str]:
    item_rows = [r for r in rows if r.kind == RowKind.ITEM]

    def where(predicate: Callable[[Item], bool]) -> str:
        return refs([r for r in item_rows if predicate(item_of(r))]) or "(none)"

    blank_expiry_ids = {r.request_id for r in records if r.kind == MessKind.BLANK_EXPIRY}
    import_day = AS_OF + timedelta(days=1)
    stale_cutoff = AS_OF - timedelta(days=30)

    def past_expiry(item: Item) -> bool:
        return (
            item.status == ItemStatus.APPROVED
            and item.expires is not None
            and item.expires < import_day
            and item.request_id not in blank_expiry_ids
        )

    def rolls_year(item: Item) -> bool:
        return any(r.on.year > item.brief.submitted.year for r in item.rounds)

    def handed_off(item: Item) -> bool:
        return len({r.reviewer.initials for r in item.rounds}) > 1

    def stale(item: Item) -> bool:
        return item.status == ItemStatus.CHANGES_REQUESTED and item.rounds[-1].on < stale_cutoff

    multi_newline = sum(1 for i in items if i.newline_comments and len(i.rounds) > 1)
    return [
        "## Other characteristics importer and metrics authors need",
        "",
        "Not errors, but each one changes what a naive importer computes.",
        "",
        f"- **Comments that cross New Year** (comment dates need the year rollover): {where(rolls_year)}.",
        (
            "- **Reviewer hand-offs** (initials change between rounds; Assigned Reviewer is the latest): "
            f"{where(handed_off)}."
        ),
        (
            "- **Version # = Revision Round + 1** (marketer sent v2 before the first review; Comments say so): "
            f"{where(lambda i: i.pre_revised)}."
        ),
        f"- **Newline-separated comment segments**: {multi_newline} multi-round cells; the rest use single spaces.",
        f"- **Five-round items**: {where(lambda i: len(i.rounds) >= MAX_ROUNDS)}.",
        (
            f"- **Stale Changes Requested** (last round before {stale_cutoff.isoformat()}, marketer never "
            f"resubmitted): {where(stale)}."
        ),
        f"- **Rejected**: {where(lambda i: i.status == ItemStatus.REJECTED)}.",
        (
            "- **Re-review requests** (new row for an expiring or rate-changed placement; the original row still says "
            "Approved; the asset name ends `orig MCR-..`, so an importer may link them): "
            f"{where(lambda i: i.brief.re_review_of is not None)}."
        ),
        (
            f"- **Approved, expiry already past on {import_day.isoformat()}**: {where(past_expiry)}. Legacy items "
            "are excluded from the expiry sweep (spec 4.2.7), so these import as approved with a past `expires_at`."
        ),
        (
            "- **Rate sheet change 6/1/26**: the PL APR floor moved from 8.99% to 9.49%. Comments before 6/1 quote "
            "8.99%; some later ones flag the old floor."
        ),
        "- **Reg B rewrite 7/21/26**: fair-lending comments after that date cite the rewritten 1002.4(b).",
        "- **Holiday freeze 12/22/25-1/2/26**: no comment segment is dated inside it (the merged note row says so).",
        "- Partner rows always postdate that partner's agreement date on the Partners sheet.",
        "",
    ]


def totals_section(rows: Sequence[SheetRow], items: Sequence[Item]) -> list[str]:
    status_counts = Counter(canonical_status(i) for i in items)
    open_items = [i for i in items if i.status in {ItemStatus.SUBMITTED, ItemStatus.IN_REVIEW}]
    lines = [
        "## Expected import totals",
        "",
        f"- Sheet rows below the header: {len(rows)} (r2-r{rows[-1].excel_row}).",
        (
            f"- Data rows with a Request ID: {len(data_rows(rows))}; note rows: "
            f"{sum(1 for r in rows if r.kind == RowKind.NOTE)}; exact duplicates to drop: "
            f"{sum(1 for r in rows if r.kind == RowKind.DUPLICATE)}."
        ),
        (
            f"- **Items after import: {len(items)}**, with {len({i.request_id for i in items})} distinct Request "
            "IDs (one ID is shared by two items)."
        ),
        (
            f"- Submitted {items[0].brief.submitted.isoformat()} to {items[-1].brief.submitted.isoformat()}; the "
            f"tracker was last updated {AS_OF.isoformat()}."
        ),
        (
            f"- Open (submitted or in review) on {AS_OF.isoformat()}: {len(open_items)}; unassigned: "
            f"{sum(1 for i in items if i.reviewer is None)}."
        ),
        "",
        "| Canonical status | Items |",
        "|---|---|",
    ]
    lines.extend(f"| {status} | {count} |" for status, count in sorted(status_counts.items()))
    lines.append("")
    return lines


def vocabulary_section() -> list[str]:
    def pairs(mapping: Sequence[tuple[str, str]]) -> str:
        return ", ".join(f"`{label}` -> `{value}`" for label, value in mapping)

    status_pairs = [
        (label, f"{status}" + (f" + decision {outcome}" if outcome is not None else ""))
        for (status, outcome), label in STATUS_DEFAULT_LABEL.items()
    ]
    return [
        "### Clean labels to app vocabulary (`backend/core/domain.py`)",
        "",
        f"- Product: {pairs([(PRODUCT_LABEL[p], str(p)) for p in Product])}.",
        f"- Channel: {pairs([(str(c), c.name.lower()) for c in Channel])}.",
        f"- Risk Tier: {pairs([(str(r), r.name.lower()) for r in Risk])}.",
        f"- Status: {pairs(status_pairs)}.",
        f"- Partners sheet Type: {pairs([(label, str(kind)) for kind, label in PARTNER_KIND_LABEL.items()])}.",
        "",
    ]


def build_readme(rows: list[SheetRow], items: list[Item], records: list[MessRecord]) -> str:
    def partner_name(item: Item) -> str:
        return item.brief.partner.name if item.brief.partner is not None else "(none)"

    tables = (
        ValueTable("Status values", Col.STATUS, canonical_status, frozenset(display(v) for v in STATUS_LIST)),
        ValueTable(
            "Product values",
            Col.PRODUCT,
            lambda i: str(i.brief.product),
            frozenset(display(v) for v in PRODUCT_LABEL.values()),
        ),
        ValueTable(
            "Trigger Terms values",
            Col.TRIGGER,
            lambda i: "Y" if i.trigger_terms else "N",
            frozenset({display("Y"), display("N")}),
        ),
        ValueTable(
            "Partner values",
            Col.PARTNER,
            partner_name,
            frozenset({display(None), *(display(p.name) for p in PARTNERS)}),
        ),
    )
    lines = [
        README_INTRO,
        "## Sheets\n",
        md_table(("Sheet", "Contents"), SHEET_ROWS),
        "## Column semantics (clean case)\n",
        md_table(("Column", "Meaning when clean"), COLUMN_ROWS),
        COMMENT_GRAMMAR,
    ]
    lines.extend(totals_section(rows, items))
    lines.extend(mess_section(records))
    lines.extend(link_section(rows))
    lines.extend(
        [
            "## Every raw value, mapped",
            "",
            (
                "Exhaustive over all data rows (the duplicate row included). For Trigger Terms the Canonical column "
                "is the generator's truth; the importer must still store blank and `?` as null."
            ),
            "",
        ]
    )
    for table in tables:
        lines.extend(value_table(rows, table))
    lines.extend(vocabulary_section())
    lines.extend(characteristics_section(rows, items, records))
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------------------------


def main() -> None:
    rng = random.Random(SEED)  # noqa: S311 - deterministic sample data, not security
    items = Generator(rng).build()
    rows = layout_rows(items, rng)
    records = apply_mess(rows, items, rng)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_workbook(rows).save(XLSX_PATH)
    README_PATH.write_text(build_readme(rows, items, records), encoding="utf-8")
    sys.stdout.write(f"wrote {XLSX_PATH} ({len(rows)} rows) and {README_PATH}\n")


if __name__ == "__main__":
    main()
